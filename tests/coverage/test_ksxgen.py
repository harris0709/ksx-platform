"""ksxgen 회귀 시험: 커버리지·생성물 최신 여부·Turtle 문법. pytest 또는 python으로 실행."""
import os
import subprocess
import sys
import tempfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "tools", "ksxgen"))

from ksxgen import cli, emit, vocab  # noqa: E402
from ksxgen import ext as X  # noqa: E402
from ksxgen.model import load  # noqa: E402

STD = os.path.join(ROOT, "standards")
EXT = os.path.join(ROOT, "extensions", "impix")


def test_counts_and_coverage():
    m = load(STD)
    v = vocab.build(m)
    assert len(m.tables) == 76 and m.noun_count == 68
    assert sum(t.role == "line" for t in m.tables.values()) == 8
    assert sum(len(x) for x in v.refs.values()) + len(v.lines) == sum(r.physical for r in m.relations) == 609
    _, ok = emit.coverage(m, v)
    assert ok and not v.issues


def test_generated_files_are_current():
    with tempfile.TemporaryDirectory() as tmp:
        cli.build(STD, tmp, cli._cfg(os.path.join(ROOT, "ksxgen.yaml")), quiet=True, exts=[EXT])
        new, cur = cli._digest(tmp), cli._digest(os.path.join(ROOT, "ontologies"))
        stale = [k for k in new if cur.get(k) != new[k]]
    assert not stale, f"ksxgen build 다시 실행 필요: {stale[:5]}"


def test_turtle_parses():
    try:
        import rdflib
    except ImportError:
        print("rdflib 없음 — Turtle 문법 시험 생략")
        return
    for f in ("tbox/ksx9101_core.ttl", "shapes/ksx9101_conformance.ttl", "vocab/codelist.ttl"):
        g = rdflib.Graph()
        g.parse(os.path.join(ROOT, "ontologies", f), format="turtle")
        assert len(g) > 1000, f


def test_extension_impix():
    """IMPIX 확장: 오류 0, 확장 명사 5·확장 속성 8, Core 이름 추가 없음."""
    m = load(STD)
    v = vocab.build(m)
    cfg = X.ext_config(EXT)
    e = X.read(cfg["path"], m, cfg)
    X.complete(e, m)
    issues = X.validate(e, m)
    assert not [i for i in issues if i.level == "오류"], issues
    kinds = [t.ext_type for t in e.tables.values()]
    assert kinds.count("확장 명사") == 5 and kinds.count("확장 속성") == 8
    props, _, _, viss = X.ext_vocab(e, m, v)
    assert not viss
    assert all(p.name.startswith("impix:") for p in props.values() if p.new)
    ttl = open(os.path.join(ROOT, "ontologies", "tbox", "impix_ext.ttl"), encoding="utf-8").read()
    assert not [l for l in ttl.splitlines() if l.startswith("ksx:")], "확장 T-Box에 ksx: 이름 선언 금지"
    assert "owl:equivalentClass" not in ttl and "owl:disjointWith" not in ttl and "rdfs:domain" not in ttl


def test_extension_prohibitions():
    """금지 규칙: Core 테이블명·속성 타입 재정의, 고정형 표준 그룹 코드 추가는 오류."""
    import openpyxl
    import shutil
    with tempfile.TemporaryDirectory() as tmp:
        d = os.path.join(tmp, "bad")
        shutil.copytree(EXT, d)
        p = os.path.join(d, "확장정의서.xlsx")
        wb = openpyxl.load_workbook(p)
        wb["테이블목록"].append([None, "x", "확장", None, "로트", "Lot", "Lot", "x", None, "EnterpriseID + LotID", None, None,
                                 None, None, None, "확장 명사", None, None, "x", None, "초안"])
        wb["테이블정의서"].append(["CustomerComplaint", "x", None, None, "x", "LotCreatedQuantity", None, "확장", "string",
                                  None, None, None, None, None, None, None, None, "x", "x", None, "초안"])
        wb["코드그룹목록"].append([None, "x", "EquipmentState", "x", "고정형", None, "x", "x", "x", "표준 그룹 하위 확장", "초안"])
        wb.save(p)
        m = load(STD)
        e = X.read(p, m, X.ext_config(d))
        X.complete(e, m)
        msgs = " | ".join(i.msg for i in X.validate(e, m) if i.level == "오류")
    assert "표준 테이블명" in msgs and "재정의 금지" in msgs and "고정형 표준 그룹" in msgs, msgs


def test_extension_turtle_parses():
    try:
        import rdflib
    except ImportError:
        print("rdflib 없음 — 확장 Turtle 문법 시험 생략")
        return
    for f in ("tbox/impix_ext.ttl", "shapes/impix_ext_shapes.ttl", "vocab/impix_codelist.ttl"):
        g = rdflib.Graph()
        g.parse(os.path.join(ROOT, "ontologies", f), format="turtle")
        assert len(g) > 100, f


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("PASS", name)
