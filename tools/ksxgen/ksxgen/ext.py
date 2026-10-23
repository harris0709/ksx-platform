"""기업 확장(Extension) — 확장 정의서 읽기·채움(sync)·검사·생성(구현안내서 5.6, 7장).

Core(standards/)는 읽기만 한다. 확장 이름은 기업 네임스페이스에 만들고, 표준 속성과 같은 영문명은
같은 의미·타입일 때 표준 공유 속성(ksx:…)을 그대로 쓴다. 금지 규칙을 어기면 오류로 보고하고 빌드를 멈춘다.
"""
from __future__ import annotations

import os
from collections import OrderedDict, defaultdict
from dataclasses import dataclass, field

import openpyxl

from . import extform as F
from .model import Code, CodeGroup, Column, Model, Relation, Table, _rows, _s
from .naming import XSD, alias, iri_template, lower_camel, strip_id
from .vocab import Vocab
from .emit import PREFIXES, lit

PK_DESC = ("— 테이블 행을 식별하는 일련번호(PK, 대리키) — 표준 외 물리 컬럼\n"
           "— 자동 증가 값(IDENTITY)으로 부여하며 다른 테이블이 참조하지 않음(FK는 UK 참조)")
BASE_SET = set(F.BASE_ROWS)


# ================================================================== 설정·데이터
@dataclass
class ExtTable:
    t: Table
    ext_type: str              # 확장 명사 / 확장 속성
    target: str | None         # 확장 속성의 대상 Core 명사
    rule: str = ""
    source: str = ""
    refs: str = ""
    state: str = ""
    uk_basis: str = ""
    uk_note: str = ""
    extra: dict = field(default_factory=dict)   # 열 단위 확장 열(원천·근거·상태): col.en → (원천, 근거, 상태)


@dataclass
class Ext:
    ent: str
    prefix: str
    ns: str
    schema: str
    path: str
    tables: "OrderedDict[str, ExtTable]"
    relations: list[Relation]
    groups: dict[str, CodeGroup]
    group_meta: dict[str, dict]
    codes: list[Code]
    code_meta: list[dict]
    judgements: list[list]
    approvals: list[list]
    orphans: list[str] = field(default_factory=list)


def ext_config(ext_dir: str) -> dict:
    cfg = {"file": "확장정의서.xlsx", "schema": "ext"}
    p = os.path.join(ext_dir, "ext.yaml")
    if os.path.exists(p):
        for line in open(p, encoding="utf-8"):
            line = line.split("#", 1)[0].strip()
            if ":" in line:
                k, v = line.split(":", 1)
                cfg[k.strip()] = v.strip().strip('"').strip("'")
    ent = cfg.get("ent") or os.path.basename(os.path.normpath(ext_dir))
    cfg.setdefault("ent", ent)
    cfg.setdefault("prefix", ent.lower())
    cfg.setdefault("ns", f"https://ontology.impix.co.kr/ext/{ent.lower()}/1.0/")
    cfg["path"] = os.path.join(ext_dir, cfg["file"])
    return cfg


def _yn(v) -> bool:
    return _s(v).upper() in ("Y", "UK")


def _num(v):
    return int(v) if isinstance(v, (int, float)) else None


def read(path: str, core: Model, cfg: dict | None = None) -> Ext:
    """확장 정의서를 읽는다(예시 행 제외). 파생 행은 sync()가 채운다."""
    cfg = cfg or {"ent": "", "prefix": "ext", "ns": "urn:ext:", "schema": "ext"}
    wb = openpyxl.load_workbook(path, data_only=True)
    tables: OrderedDict[str, ExtTable] = OrderedDict()
    for r in _rows(wb["테이블목록"]):
        if _s(r.get("상태")) == "예시":
            continue
        etype, target = _s(r.get("확장유형")), _s(r.get("대상 Core 명사")) or None
        en = _s(r.get("테이블명(영문)")) or ((target + "Ext") if etype == "확장 속성" and target else "")
        if not en:
            continue
        uk = [x.strip() for x in _s(r.get("UK(식별 키)")).split("+") if x.strip()]
        t = Table(en=en, ko=_s(r.get("테이블명(국문)")), area=_s(r.get("업무영역")), clause="확장", table_no=None,
                  definition=_s(r.get("테이블 설명(표준 정의)")), uk=uk)
        tables[en] = ExtTable(t, etype, target, _s(r.get("판정 규칙(5.6)")), _s(r.get("원천(시스템.테이블)")),
                              _s(r.get("관련 시나리오·KPI·질의")), _s(r.get("상태")), _s(r.get("UK 근거")),
                              _s(r.get("UK 비고")))
    orphans: list[str] = []
    for r in _rows(wb["테이블정의서"]):
        if _s(r.get("상태")) == "예시":
            continue
        et = tables.get(_s(r.get("테이블명(영문)")))
        en = _s(r.get("컬럼명(영문)"))
        if et is None and en:
            orphans.append(f"{_s(r.get('테이블명(영문)'))}.{en}")
        if et is None or not en:
            continue
        et.t.columns.append(Column(
            table=et.t.en, seq=_num(r.get("순번")) or 999, ko=_s(r.get("속성명(국문)")), en=en, kind=_s(r.get("구분")) or "확장",
            dtype=_s(r.get("데이터 타입(표준)")), ptype=_s(r.get("물리 타입(권장)")), pk=_yn(r.get("PK")),
            uk=_s(r.get("UK")) == "UK", fk=_yn(r.get("FK")), notnull=_yn(r.get("NOT NULL")),
            fkref=_s(r.get("FK 참조(테이블.컬럼)")) or None, ref=_s(r.get("참조속성대상여부")) or None,
            code=_s(r.get("코드 참조")) or None, desc=_s(r.get("설명"))))
        et.extra[en] = (_s(r.get("원천(시스템.테이블.컬럼)")), _s(r.get("근거(시나리오·KPI·질의)")), _s(r.get("상태")))
    groups, gmeta = {}, {}
    for r in _rows(wb["코드그룹목록"]):
        g = _s(r.get("공통코드그룹(CommonCodeGroup)"))
        if not g or _s(r.get("상태")) == "예시":
            continue
        groups[g] = CodeGroup(g, _s(r.get("그룹명")), _s(r.get("업무영역")), _s(r.get("코드체계")) or None,
                              _s(r.get("상위 그룹")) or None, _s(r.get("그룹 설명")) or None)
        gmeta[g] = {"kind": _s(r.get("확장 구분")) or "신규 그룹", "attrs": _s(r.get("사용 속성(영문)")),
                    "nouns": _s(r.get("사용 명사")), "state": _s(r.get("상태"))}
    codes, cmeta = [], []
    for r in _rows(wb["공통코드"]):
        g, c = _s(r.get("공통코드그룹(CommonCodeGroup)")), _s(r.get("공통코드식별자(CommonCodeID)"))
        if not (g and c) or _s(r.get("상태")) == "예시":
            continue
        codes.append(Code(g, c, _s(r.get("공통코드명(CommonCodeName)")), _s(r.get("영문명")) or None,
                          _num(r.get("표시순번(CommonCodeSequence)")), _s(r.get("사용여부(UseIndicator)")) or "Y",
                          _s(r.get("상위공통코드그룹(ParentCommonCodeGroup)")) or None,
                          _s(r.get("상위공통코드식별자(ParentCommonCodeID)")) or None, _s(r.get("설명")) or None))
        cmeta.append({"area": _s(r.get("업무영역")), "group_name": _s(r.get("그룹명")), "src": _s(r.get("출처구분")) or "확장",
                      "apply": _s(r.get("회사 적용(Y/N)")) or "Y", "cval": _s(r.get("회사 코드값")),
                      "cname": _s(r.get("회사 코드명")), "note": _s(r.get("비고(회사)")), "state": _s(r.get("상태"))})
    judg = [[r.get(h.replace("\n", "")) for h in F.JUDGE] for r in _rows(wb["판정목록"]) if _s(r.get("상태")) != "예시"
            and any(r.get(h) for h in ("원천(시스템.테이블.컬럼)", "판정"))]
    appr = [[r.get(h) for h in F.APPROVE] for r in _rows(wb["검토·승인"]) if _s(r.get("구분")) != "검사"]
    return Ext(cfg.get("ent", ""), cfg.get("prefix", "ext"), cfg.get("ns", ""), cfg.get("schema", "ext"), path,
               tables, [], groups, gmeta, codes, cmeta, judg, appr, orphans)


# ================================================================== 채움(sync)
def _core_col(core: Model, table: str, en: str) -> Column | None:
    t = core.tables.get(table)
    return t.col(en) if t else None


def complete(ext: Ext, core: Model) -> None:
    """작성자가 비워 둔 파생 정보를 채운다: PK·기본 속성·Core키 행, UK, 물리명·타입, 참조속성 대상여부, 관계."""
    base_src = core.tables["Lot"]
    for et in ext.tables.values():
        t = et.t
        user = [c for c in t.columns if c.kind not in ("PK", "기본", "Core키") and not c.pk]
        given = {c.en: c for c in t.columns}
        cols: list[Column] = [Column(t.en, 0, f"{t.ko} 일련번호", f"{t.en}SequenceNo", "PK", "bigint", "BIGINT", True, False,
                                     False, True, None, "비대상", None, PK_DESC)]
        if et.ext_type == "확장 속성" and et.target in core.tables:
            ct = core.tables[et.target]
            if not t.uk:
                t.uk = list(ct.uk)
            for k in ct.uk:
                cc = ct.col(k)
                cols.append(Column(t.en, 0, cc.ko if cc else k, k, "Core키", cc.dtype if cc else "IDREF",
                                   cc.ptype if cc else "VARCHAR(50)", False, True, True, True, f"{ct.en}.{k}", "대상", None,
                                   f"— 확장 대상 Core 명사 {ct.ko}({ct.en})의 식별 키(UK) — 1:1 확장\n— 참조 대상: {ct.ko}({ct.en})"))
        elif et.ext_type == "확장 명사":
            for b in F.BASE_ROWS:
                bc = given.get(b) or base_src.col(b)
                cols.append(Column(t.en, 0, bc.ko, b, "기본", bc.dtype, bc.ptype, False, False, bc.fk, bc.notnull,
                                   bc.fkref, bc.ref, bc.code, bc.desc))
            if not t.uk:
                idc = next((c for c in user if c.dtype == "ID"), None)
                t.uk = ["EnterpriseID"] + ([idc.en] if idc else [])
        for c in user:
            c.kind = "확장"
            c.ptype = c.ptype or F.PTYPE.get(c.dtype, "VARCHAR(255)")
            c.uk = c.en in t.uk
            c.notnull = c.notnull or c.uk
            c.fk = bool(c.fkref)
            c.ref = c.ref or ("대상" if c.dtype in ("IDREF", "enum") else "비대상")
            cols.append(c)
        for c in cols:
            if c.kind == "기본" and c.en == "EnterpriseID":
                c.uk = True
        for i, c in enumerate(cols, 1):
            c.seq = i
        t.columns = cols
    ext.relations = derive_relations(ext, core)


def _lookup(ext: Ext, core: Model, name: str) -> Table | None:
    return core.tables.get(name) or (ext.tables[name].t if name in ext.tables else None)


def derive_relations(ext: Ext, core: Model) -> list[Relation]:
    """FK 참조 열 → 관계. 부모 UK가 복합(문서번호+순번 등)이면 같은 부모를 가리키는 연속 컬럼을 하나로 묶는다."""
    rels: list[Relation] = []
    for et in ext.tables.values():
        t = et.t
        groups: list[tuple[str, list[Column], list[str]]] = []
        core_key: list = []
        for c in t.columns:
            if not c.fkref or "." not in c.fkref:
                continue
            pt, pc = c.fkref.split(".", 1)
            if c.kind == "Core키":
                core_key.append((c, pt, pc))
                continue
            parent = _lookup(ext, core, pt)
            width = len([k for k in parent.uk if k != "EnterpriseID"]) if parent else 1
            if groups and groups[-1][0] == pt and len(groups[-1][1]) < width and c.en != "EnterpriseID":
                groups[-1][1].append(c)
                groups[-1][2].append(pc)
            else:
                groups.append((pt, [c], [pc]))
        if core_key:
            groups.insert(0, (core_key[0][1], [x[0] for x in core_key], [x[2] for x in core_key]))
        for pt, cols, pcols in groups:
            fk = [c.en for c in cols]
            if fk[0] != "EnterpriseID":
                fk, pcols = ["EnterpriseID"] + fk, ["EnterpriseID"] + pcols
            kind = ("확장 대상(1:1)" if cols[0].kind == "Core키" else
                    "기본 속성" if cols[0].en in BASE_SET else "자기 참조" if pt == t.en else "명사 참조")
            rels.append(Relation(len(rels) + 1, t.en, fk, pt, pcols, kind, None, True))
    return rels


# ================================================================== 검사
@dataclass
class Issue:
    level: str        # 오류 / 경고
    where: str
    msg: str


def _core_attr_types(core: Model) -> dict[str, set[str]]:
    d = defaultdict(set)
    for t in core.tables.values():
        for c in t.columns:
            if not c.pk:
                d[c.en].add(c.dtype)
    return d


def _compat(a: str, b: str) -> bool:
    ids = {"ID", "IDREF", "string"}
    return a == b or (a in ("ID", "IDREF") and b in ("ID", "IDREF")) or (a in ids and b in ids and "string" not in (a, b))


def validate(ext: Ext, core: Model) -> list[Issue]:
    out: list[Issue] = []
    E = lambda w, m: out.append(Issue("오류", w, m))       # noqa: E731
    W = lambda w, m: out.append(Issue("경고", w, m))       # noqa: E731
    ctypes = _core_attr_types(core)
    for o in ext.orphans:
        E(o, "테이블정의서 행의 테이블이 테이블목록에 없음 — 테이블목록에 먼저 등록")
    all_groups = set(core.code_groups) | set(ext.groups)
    for et in ext.tables.values():
        t, w = et.t, et.t.en
        if et.ext_type not in F.EXT_TYPES:
            E(w, f"확장유형 '{et.ext_type}'은 허용되지 않음({', '.join(F.EXT_TYPES)})")
        if t.en.lower().startswith("ksx"):
            E(w, "ksx 접두 이름 금지(이름 공간은 기업 prefix)")
        if et.ext_type == "확장 명사":
            if t.en in core.tables:
                E(w, "표준 테이블명과 같은 이름 — Core 재정의 금지(확장 속성으로 판정하거나 이름 변경)")
            if not t.uk or t.uk[0] != "EnterpriseID" or len(t.uk) < 2:
                E(w, "UK는 EnterpriseID + 식별 속성(ID 또는 조합)이어야 함")
            if not any(c.dtype == "ID" and c.uk for c in t.columns) and "조합" not in et.uk_basis:
                W(w, "ID 식별 속성이 없음 — 조합 식별이면 UK 근거에 '조합'과 이유를 적는다")
            for b in ("EnterpriseID", "SiteID", "PlantID"):
                if not t.col(b):
                    E(w, f"기본 속성 {b} 없음(ext-sync로 채움)")
        elif et.ext_type == "확장 속성":
            if et.target not in core.tables:
                E(w, f"대상 Core 명사 '{et.target}'가 표준에 없음")
                continue
            ct = core.tables[et.target]
            if t.en != ct.en + "Ext":
                E(w, f"확장 속성 테이블명은 {ct.en}Ext")
            if t.uk != ct.uk:
                E(w, f"UK는 대상 Core UK와 같아야 함({' + '.join(ct.uk)})")
        for c in t.columns:
            cw = f"{w}.{c.en}"
            if c.kind in ("PK", "Core키") or (c.kind == "기본" and c.en in BASE_SET):
                continue
            if c.dtype not in F.DTYPES:
                E(cw, f"데이터 타입 '{c.dtype}' 허용 안 됨({', '.join(F.DTYPES)})")
            if et.ext_type == "확장 속성" and et.target in core.tables and core.tables[et.target].col(c.en):
                E(cw, f"표준 명사 {et.target}에 이미 있는 속성 — 확장이 아니라 표준 매핑")
            if c.en in ctypes and not any(_compat(c.dtype, x) for x in ctypes[c.en]):
                E(cw, f"표준 속성과 같은 이름에 다른 타입 — 재정의 금지(표준 {sorted(ctypes[c.en])}, 확장 {c.dtype})")
            if c.dtype == "IDREF":
                if not c.fkref or "." not in c.fkref:
                    E(cw, "IDREF는 FK 참조(테이블.컬럼) 필요")
                else:
                    pt, pc = c.fkref.split(".", 1)
                    p = _lookup(ext, core, pt)
                    if p is None:
                        E(cw, f"FK 참조 테이블 '{pt}' 없음(표준·확장 모두)")
                    elif pc not in p.uk:
                        E(cw, f"FK 참조 컬럼 {pt}.{pc}가 부모 UK({' + '.join(p.uk)})에 없음")
            if c.dtype == "enum":
                g = c.code_group
                if not c.code:
                    E(cw, "enum은 코드 참조 필요(CommonCodeGroup=… 또는 단위그룹(UOMGroup).UOMCode)")
                elif g and g not in all_groups:
                    E(cw, f"코드 그룹 '{g}'가 표준·확장 어디에도 없음")
            if not c.desc:
                W(cw, "설명 없음")
            src, _, st = et.extra.get(c.en, ("", "", ""))
            if not src and st != "후보":
                W(cw, "원천 없음 — 원천 컬럼을 적거나 상태를 '후보'로")
    seen = set()
    for g, meta in ext.group_meta.items():
        if meta["kind"] == "신규 그룹" and g in core.code_groups:
            E(f"코드그룹 {g}", "표준 그룹과 같은 이름 — '표준 그룹 하위 확장'으로 적거나 이름 변경")
        if meta["kind"] == "표준 그룹 하위 확장":
            if g not in core.code_groups:
                E(f"코드그룹 {g}", "하위 확장 대상 표준 그룹이 없음")
            elif "회사정의" not in (core.code_groups[g].scheme or ""):
                E(f"코드그룹 {g}", f"고정형 표준 그룹({core.code_groups[g].scheme})에는 코드 추가 금지 — 코드 대응으로 처리")
    core_codes = {(c.group, c.code) for c in core.codes}
    ext_codes = {(c.group, c.code) for c in ext.codes}
    for c in ext.codes:
        where = f"코드 {c.group}/{c.code}"
        if c.group not in ext.groups:
            E(where, "코드그룹목록에 그룹이 없음")
        if (c.group, c.code) in core_codes:
            E(where, "표준 코드와 같은 코드 — 재정의 금지(코드 대응으로 처리)")
        if (c.group, c.code) in seen:
            E(where, "중복 코드")
        seen.add((c.group, c.code))
        if c.parent_group and c.parent_code and (c.parent_group, c.parent_code) not in core_codes | ext_codes:
            E(where, f"상위 코드 {c.parent_group}/{c.parent_code} 없음")
    return out


# ================================================================== 시트 쓰기(sync)
def _ddl_table(ext: Ext, core: Model, et: ExtTable) -> str:
    t, s = et.t, ext.schema
    lines = []
    for c in t.columns:
        if c.pk:
            lines.append(f'    "{c.en}" BIGINT GENERATED BY DEFAULT AS IDENTITY, -- {c.ko}(PK)')
        else:
            lines.append(f'    "{c.en}" {c.ptype or "VARCHAR(255)"}' + (" NOT NULL" if c.notnull else "") + f", -- {c.ko}")
    lines.append(f'    CONSTRAINT "PK_{t.en}" PRIMARY KEY ("{t.en}SequenceNo"),')
    lines.append(f'    CONSTRAINT "UK_{t.en}" UNIQUE (' + ", ".join(f'"{k}"' for k in t.uk) + ")")
    head = f"확장 속성 → {et.target}" if et.ext_type == "확장 속성" else "확장 명사"
    return f'CREATE TABLE {s}."{t.en}" ( -- {t.ko} ({head})\n' + "\n".join(lines) + "\n);"


def ddl_sql(ext: Ext, core: Model, core_schema: str = "ksx") -> str:
    s = ext.schema
    o = [f"-- {ext.ent} 기업 확장 DDL (ksxgen 생성물, 직접 수정 금지) — Core 스키마 {core_schema} 참조",
         f"CREATE SCHEMA IF NOT EXISTS {s};", ""]
    for et in ext.tables.values():
        o.append(_ddl_table(ext, core, et))
        o.append(f"COMMENT ON TABLE {s}.\"{et.t.en}\" IS '{et.t.ko} — {et.t.definition.replace(chr(39), chr(39) * 2)}';")
        for c in et.t.columns:
            o.append(f"COMMENT ON COLUMN {s}.\"{et.t.en}\".\"{c.en}\" IS '{c.ko.replace(chr(39), chr(39) * 2)}';")
        o.append("")
    n = defaultdict(int)
    o.append("-- 물리 FK: 부모 UK 참조(표준 → ksx 스키마, 확장 → ext 스키마)")
    for r in ext.relations:
        n[r.child] += 1
        ps = s if r.parent in ext.tables else core_schema
        o.append(f'ALTER TABLE {s}."{r.child}" ADD CONSTRAINT "FK_{r.child}_{n[r.child]:02d}" FOREIGN KEY ('
                 + ", ".join(f'"{c}"' for c in r.fk_cols) + f') REFERENCES {ps}."{r.parent}" ('
                 + ", ".join(f'"{c}"' for c in r.parent_cols) + ") DEFERRABLE INITIALLY IMMEDIATE;")
    return "\n".join(o) + "\n"


def write_back(ext: Ext, core: Model, issues: list[Issue], out_path: str | None = None) -> str:
    wb = F.new_workbook()
    ws = wb["테이블목록"]
    nrow_td = sum(len(et.t.columns) for et in ext.tables.values()) + 1
    nrow_rel = len(ext.relations) + 1
    for i, et in enumerate(ext.tables.values(), 1):
        t, r = et.t, i + 1
        basis = et.uk_basis or ("Core UK" if et.ext_type == "확장 속성" else
                                ("확장 ID" if any(c.dtype == "ID" and c.uk for c in t.columns) else "조합(제안)"))
        ws.append([i, t.area, "확장", None, t.ko, t.en, t.en, t.definition, f"{t.en}SequenceNo", " + ".join(t.uk), basis,
                   et.uk_note or None,
                   f"=COUNTIF(테이블정의서!$A$2:$A${nrow_td},F{r})",
                   f'=COUNTIFS(테이블정의서!$A$2:$A${nrow_td},F{r},테이블정의서!$H$2:$H${nrow_td},"기본")',
                   f'=COUNTIFS(관계목록!$B$2:$B${nrow_rel},F{r},관계목록!$H$2:$H${nrow_rel},"물리")',
                   et.ext_type, et.target, et.rule or ("② 확장 속성" if et.ext_type == "확장 속성" else "③ 확장 명사"),
                   et.source, et.refs, et.state or "초안"])
    ws = wb["테이블정의서"]
    for et in ext.tables.values():
        for c in et.t.columns:
            src, why, st = et.extra.get(c.en, ("", "", ""))
            if c.kind in ("PK", "기본", "Core키"):
                src, why, st = src, why, st or "자동"
            ws.append([et.t.en, et.t.ko, et.t.en, c.seq, c.ko, c.en, c.en, c.kind, c.dtype, c.ptype or None,
                       "Y" if c.pk else None, "UK" if c.uk else None, "Y" if c.fk else None,
                       "Y" if c.notnull else None, c.fkref, c.ref, c.code, c.desc or None, src or None, why or None,
                       st or et.state or "초안"])
    ws = wb["관계목록"]
    for r in ext.relations:
        note = "Core 행 1개당 확장 행 최대 1개" if r.kind == "확장 대상(1:1)" else None
        ws.append([r.no, r.child, ", ".join(r.fk_cols), r.parent, ", ".join(r.parent_cols), r.kind, note, "물리"])
    ws = wb["코드그룹목록"]
    for i, (g, cg) in enumerate(ext.groups.items(), 1):
        m = ext.group_meta[g]
        ws.append([i, cg.area, g, cg.name, cg.scheme, cg.parent, cg.desc, m["attrs"], m["nouns"], m["kind"], m["state"] or "초안"])
    ws = wb["공통코드"]
    for i, (c, m) in enumerate(zip(ext.codes, ext.code_meta), 1):
        ws.append([i, m["area"], c.group, m["group_name"], c.code, c.name, c.en, c.seq, c.use, c.parent_group,
                   c.parent_code, c.desc, m["src"], m["apply"], m["cval"] or None, m["cname"] or None, m["note"] or None,
                   m["state"] or "초안"])
    ws = wb["판정목록"]
    for i, j in enumerate(ext.judgements, 1):
        ws.append([i] + list(j[1:]))
    ws = wb["검토·승인"]
    ne, nw = sum(1 for x in issues if x.level == "오류"), sum(1 for x in issues if x.level == "경고")
    ws.append(["검사", "ksxgen ext-sync 결과", f"오류 {ne}, 경고 {nw} — 확장 명사 "
               f"{sum(1 for e in ext.tables.values() if e.ext_type == '확장 명사')}, 확장 속성 테이블 "
               f"{sum(1 for e in ext.tables.values() if e.ext_type == '확장 속성')}, 확장 컬럼 "
               f"{sum(1 for e in ext.tables.values() for c in e.t.columns if c.kind == '확장')}, 코드 그룹 {len(ext.groups)}, "
               f"코드 {len(ext.codes)}, 판정 {len(ext.judgements)}", None, None])
    for x in issues:
        ws.append(["검사", f"{x.level}: {x.where}", x.msg, None, None])
    for a in ext.approvals:
        ws.append(list(a))
    ws = wb["DDL"]
    for et in ext.tables.values():
        ws.append([et.t.en, _ddl_table(ext, core, et)])
    fk = [l for l in ddl_sql(ext, core).splitlines() if l.startswith("ALTER TABLE")]
    ws.append(["(물리 FK)", "\n".join(fk)])
    for name in ("테이블목록", "테이블정의서", "관계목록", "코드그룹목록", "공통코드", "판정목록", "검토·승인", "DDL"):
        F._style_body(wb[name])
    for row in wb["DDL"].iter_rows(min_row=2):
        row[1].font = openpyxl.styles.Font(name="Consolas", size=9)
    for row in wb["테이블정의서"].iter_rows(min_row=2):
        if row[7].value in ("PK", "기본", "Core키"):
            for c in row:
                c.fill = openpyxl.styles.PatternFill("solid", fgColor="F2F2F2")
    for row in wb["검토·승인"].iter_rows(min_row=2):
        if str(row[1].value or "").startswith("오류"):
            row[1].fill = openpyxl.styles.PatternFill("solid", fgColor="F8D7DA")
        elif str(row[1].value or "").startswith("경고"):
            row[1].fill = openpyxl.styles.PatternFill("solid", fgColor="FFF3CD")
    F.add_validations(wb)
    p = out_path or ext.path
    wb.save(p)
    return p


def sync(path: str, std: str, cfg: dict | None = None):
    from .model import load
    core = load(std)
    ext = read(path, core, cfg)
    complete(ext, core)
    issues = validate(ext, core)
    write_back(ext, core, issues)
    return ext, issues


# ================================================================== 생성(S4)
@dataclass
class EProp:
    name: str            # 접두 포함 QName(ksx:… 재사용 또는 {prefix}:…)
    kind: str            # data / code / ref
    range: str
    label_ko: str
    label_en: str
    code_group: str | None
    new: bool
    sources: set = field(default_factory=set)


def ext_vocab(ext: Ext, core: Model, v: Vocab):
    """확장 컬럼·관계 → 속성. 표준 공유 속성은 재사용, 없으면 기업 네임스페이스에 신규 생성."""
    by_attr: dict[str, object] = {}
    for cps in v.colprops.values():
        for cp in cps:
            by_attr.setdefault(cp.col.en, cp.prop)
    core_names = set(v.props)
    props: "OrderedDict[str, EProp]" = OrderedDict()
    colprop: dict[tuple, EProp] = {}
    issues: list[Issue] = []
    for et in ext.tables.values():
        for c in et.t.columns:
            if c.pk or c.kind == "Core키":
                continue
            kind = "code" if c.dtype == "enum" else "data"
            rng = "skos:Concept" if kind == "code" else XSD.get(c.dtype, "xsd:string")
            p = by_attr.get(c.en)
            if p is not None:
                if (p.kind, p.range) != (kind, rng):
                    issues.append(Issue("오류", f"{et.t.en}.{c.en}", f"표준 속성 ksx:{p.name}({p.kind}:{p.range})와 타입이 다름 — 재정의 금지"))
                    continue
                q = f"ksx:{p.name}"
                ep = props.get(q) or EProp(q, kind, rng, p.label_ko, p.label_en, p.code_group, False)
            else:
                local = lower_camel(c.en)
                if local in core_names:
                    issues.append(Issue("오류", f"{et.t.en}.{c.en}", f"표준 속성 이름 ksx:{local}과 충돌 — 영문명 변경"))
                    continue
                q = f"{ext.prefix}:{local}"
                ep = props.get(q) or EProp(q, kind, rng, c.ko, c.en, c.code_group, True)
            ep.sources.add(c.en)
            props[q] = ep
            colprop[(et.t.en, c.en)] = ep
    refs: dict[str, list[tuple[Relation, EProp, bool]]] = defaultdict(list)
    for r in ext.relations:
        if r.kind == "확장 대상(1:1)":
            continue
        parent_core = r.parent in core.tables
        pcls = (f"ksx:{core.tables[r.parent].cls}" if parent_core else f"{ext.prefix}:{r.parent}")
        cols = [c for c in r.fk_cols if c != "EnterpriseID"] or r.fk_cols
        base = ("ref" + strip_id(cols[0])) if cols[0].endswith("ID") else ("ref" + r.parent)
        cp = v.props.get(base)
        if cp is not None and cp.kind == "ref" and cp.range == pcls:
            q, new = f"ksx:{base}", False
        else:
            local = base if (cp is None and base not in core_names) else base + "To" + r.parent
            q, new = f"{ext.prefix}:{local}", True
        child = ext.tables[r.child].t
        ep = props.get(q) or EProp(q, "ref", pcls, (child.col(cols[0]).ko + " 참조") if child.col(cols[0]) else q,
                                   q.split(":")[1], None, new)
        ep.sources.add(cols[0])
        props[q] = ep
        notnull = all(child.col(c) and child.col(c).notnull for c in r.fk_cols)
        refs[r.child].append((r, ep, notnull))
    return props, colprop, refs, issues


def _pfx(ext: Ext, cfg: dict) -> str:
    return PREFIXES.format(ns=cfg["ns"]) + f"@prefix {ext.prefix}: <{ext.ns}> .\n"


def tbox(ext: Ext, core: Model, props, cfg) -> str:
    o = [_pfx(ext, cfg), f"<{ext.ns.rstrip('/')}> a owl:Ontology ;",
         f"    owl:imports <{cfg['ns'].rstrip('/')}> ;",
         f"    rdfs:label {lit(ext.ent + ' 기업 확장(Extension)', 'ko')} ;",
         f"    dct:source {lit(os.path.basename(ext.path))} .", "", "### 확장 명사"]
    for et in ext.tables.values():
        t = et.t
        if et.ext_type == "확장 명사":
            o += [f"{ext.prefix}:{t.en} a owl:Class ; rdfs:subClassOf ksx:Noun ;",
                  f"    rdfs:label {lit(t.ko, 'ko')}, {lit(t.en, 'en')} ;",
                  f"    skos:definition {lit(t.definition, 'ko')} ;",
                  f"    ksx:businessArea {lit(t.area, 'ko')} ; ksx:physicalTable {lit(ext.schema + '.' + t.en)} ."]
    o += ["", "### 확장 속성(표준 공유 속성은 Core 정의를 그대로 쓰고 여기서 다시 선언하지 않는다)"]
    for p in props.values():
        if not p.new:
            continue
        if p.kind == "data":
            head = f"{p.name} a owl:DatatypeProperty ; rdfs:range {p.range}"
        elif p.kind == "code":
            head = f"{p.name} a owl:ObjectProperty ; rdfs:range skos:Concept" + (
                f" ; ksx:codeGroup {lit(p.code_group)}" if p.code_group else "")
        else:
            head = f"{p.name} a owl:ObjectProperty ; rdfs:range {p.range} ; owl:inverseOf {p.name}Inverse"
        extra = [f"rdfs:label {lit(p.label_ko, 'ko')}, {lit(p.label_en, 'en')}"]
        extra += [f"ksx:sourceAttribute {lit(a)}" for a in sorted(p.sources)]
        o.append(head + " ;\n    " + " ;\n    ".join(extra) + " .")
        if p.kind == "ref":
            o.append(f"{p.name}Inverse a owl:ObjectProperty ; rdfs:label {lit(p.label_ko + '(역)', 'ko')} .")
    return "\n".join(o) + "\n"


def shapes(ext: Ext, core: Model, colprop, refs, cfg) -> str:
    o = [_pfx(ext, cfg), f"@prefix {ext.prefix}s: <{ext.ns}shapes/> .", ""]
    for et in ext.tables.values():
        t = et.t
        cls = f"ksx:{core.tables[et.target].cls}" if et.ext_type == "확장 속성" else f"{ext.prefix}:{t.en}"
        ps = []
        for c in t.columns:
            p = colprop.get((t.en, c.en))
            if p is None:
                continue
            s = f"sh:path {p.name} ; sh:maxCount 1" + (" ; sh:minCount 1" if c.notnull and et.ext_type == "확장 명사" else "")
            s += f" ; sh:datatype {p.range}" if p.kind == "data" else " ; sh:class skos:Concept"
            if p.kind == "data" and c.dtype in ("float", "int") and c.en.endswith(("Quantity", "Amount")):
                s += " ; sh:minInclusive 0"
            ps.append(s)
        for r, p, nn in refs.get(t.en, []):
            ps.append(f"sh:path {p.name} ; sh:class {p.range}" + (" ; sh:minCount 1" if nn and et.ext_type == "확장 명사" else ""))
        body = " ;\n    ".join(f"sh:property [ {s} ]" for s in ps)
        o.append(f"{ext.prefix}s:{t.en}Shape a sh:NodeShape ; sh:targetClass {cls} ;\n    {body} .\n")
    return "\n".join(o)


def obda(ext: Ext, core: Model, colprop, refs, cfg) -> str:
    o = ["[PrefixDeclaration]", f"ksx:\t{cfg['ns']}", f"{ext.prefix}:\t{ext.ns}", f"d:\t{cfg['data_ns']}",
         "xsd:\thttp://www.w3.org/2001/XMLSchema#", "skos:\thttp://www.w3.org/2004/02/skos/core#", "",
         "[MappingDeclaration] @collection [["]
    for et in ext.tables.values():
        t = et.t
        cols = [c for c in t.columns if not c.pk]
        src = f'SELECT {", ".join(chr(34) + c.en + chr(34) + " AS " + alias(c.en) for c in cols)} FROM {ext.schema}."{t.en}"'
        if et.ext_type == "확장 속성":
            me, trip = iri_template(core.tables[et.target], core), []
        else:
            me = iri_template(t, core)
            trip = [f"{me} a {ext.prefix}:{t.en}"]
        for c in cols:
            p = colprop.get((t.en, c.en))
            if p is None:
                continue
            a = "{" + alias(c.en) + "}"
            if p.kind == "code":
                trip.append(f"{p.name} " + (f"d:codelist/{p.code_group}/{a}" if p.code_group else
                                            (f"d:uom/{a}" if c.en.endswith("UOMCode") else a)))
            elif p.range == "xsd:string":
                trip.append(f"{p.name} {a}")
            else:
                trip.append(f"{p.name} {a}^^{p.range}")
        for r, p, _ in refs.get(t.en, []):
            parent = core.tables.get(r.parent) or ext.tables[r.parent].t
            trip.append(f"{p.name} {iri_template(parent, core, dict(zip(r.parent_cols, r.fk_cols)))}")
        if not trip:
            continue
        o += [f"mappingId\t{ext.ent}_{t.en}", "target\t\t" + (me + " " if et.ext_type == "확장 속성" else "")
              + " ; ".join(trip) + " .", f"source\t\t{src}", ""]
    o.append("]]")
    return "\n".join(o) + "\n"


def codelist(ext: Ext, cfg) -> str:
    dn = cfg["data_ns"]
    o = [_pfx(ext, cfg), ""]
    for g, cg in ext.groups.items():
        if ext.group_meta[g]["kind"] == "신규 그룹":
            o.append(f"<{dn}codelist/{g}> a skos:ConceptScheme ; skos:prefLabel {lit(cg.name, 'ko')} ; "
                     f"skos:notation {lit(g)} ; skos:definition {lit(cg.desc or '', 'ko')} .")
    for c in ext.codes:
        s = (f"<{dn}codelist/{c.group}/{c.code}> a skos:Concept ; skos:inScheme <{dn}codelist/{c.group}> ; "
             f"skos:notation {lit(c.code)} ; skos:prefLabel {lit(c.name, 'ko')}" + (f", {lit(c.en, 'en')}" if c.en else ""))
        if c.parent_group and c.parent_code:
            s += f" ; skos:broader <{dn}codelist/{c.parent_group}/{c.parent_code}>"
        if c.desc:
            s += f" ; skos:definition {lit(c.desc, 'ko')}"
        o.append(s + " .")
    return "\n".join(o) + "\n"


def report(ext: Ext, core: Model, props, colprop, issues) -> str:
    nouns = [e for e in ext.tables.values() if e.ext_type == "확장 명사"]
    attrs = [e for e in ext.tables.values() if e.ext_type == "확장 속성"]
    newp = [p for p in props.values() if p.new]
    md = [f"# {ext.ent} 기업 확장 생성 리포트 (ksxgen)", "",
          "| 항목 | 값 |", "|---|---|",
          f"| 확장 명사(클래스) | {len(nouns)} — {', '.join(e.t.en for e in nouns)} |",
          f"| 확장 속성 테이블(대상 Core 명사) | {len(attrs)} — {', '.join(e.target for e in attrs)} |",
          f"| 확장 컬럼 | {sum(1 for e in ext.tables.values() for c in e.t.columns if c.kind == '확장')} |",
          f"| 속성: 표준 공유 재사용 / 신규 | {sum(1 for p in props.values() if not p.new)} / {len(newp)} |",
          f"| 관계(물리 FK) | {len(ext.relations)} |",
          f"| 확장 코드 그룹 / 코드 | {len(ext.groups)} / {len(ext.codes)} |",
          f"| 판정 기록 | {len(ext.judgements)} |",
          f"| 오류 / 경고 | {sum(1 for x in issues if x.level == '오류')} / {sum(1 for x in issues if x.level == '경고')} |", "",
          "## 속성 대응(확장 컬럼 → 온톨로지 속성)", "", "| 테이블 | 컬럼 | 속성 | 구분 |", "|---|---|---|---|"]
    for (tb, col), p in colprop.items():
        if ext.tables[tb].t.col(col).kind == "확장":
            md.append(f"| {tb} | {col} | `{p.name}` | {'신규' if p.new else '표준 재사용'} |")
    md += ["", "## 검사 결과", ""] + ([f"- {x.level} · {x.where}: {x.msg}" for x in issues] or ["- 오류·경고 없음"])
    return "\n".join(md) + "\n"


def build(ext_dir: str, core: Model, v: Vocab, cfg: dict, out: str, quiet=False) -> bool:
    from .emit import _w
    ecfg = ext_config(ext_dir)
    ext = read(ecfg["path"], core, ecfg)
    complete(ext, core)
    issues = validate(ext, core)
    props, colprop, refs, viss = ext_vocab(ext, core, v)
    seen = {x.where for x in issues if x.level == "오류"}
    issues += [x for x in viss if x.where not in seen]
    errors = [x for x in issues if x.level == "오류"]
    ent = ext.ent.lower()
    _w(os.path.join(out, "report", f"{ent}_ext.md"), report(ext, core, props, colprop, issues))
    if errors:
        if not quiet:
            print(f"ksxgen --ext {ext_dir}: 오류 {len(errors)}건 — 빌드 중단(report/{ent}_ext.md)")
            for x in errors[:20]:
                print(f"  오류 · {x.where}: {x.msg}")
        return False
    _w(os.path.join(out, "tbox", f"{ent}_ext.ttl"), tbox(ext, core, props, cfg))
    _w(os.path.join(out, "shapes", f"{ent}_ext_shapes.ttl"), shapes(ext, core, colprop, refs, cfg))
    _w(os.path.join(out, "mappings", f"{ent}_ext.obda"), obda(ext, core, colprop, refs, cfg))
    _w(os.path.join(out, "ddl", f"ext_{ent}.sql"), ddl_sql(ext, core, cfg.get("schema", "ksx")))
    _w(os.path.join(out, "vocab", f"{ent}_codelist.ttl"), codelist(ext, cfg))
    if not quiet:
        print(f"ksxgen --ext {ext_dir}: 확장 명사 {sum(1 for e in ext.tables.values() if e.ext_type == '확장 명사')} · "
              f"확장 속성 테이블 {sum(1 for e in ext.tables.values() if e.ext_type == '확장 속성')} · "
              f"속성 신규 {sum(1 for p in props.values() if p.new)}/재사용 {sum(1 for p in props.values() if not p.new)} · "
              f"코드 {len(ext.codes)} · 경고 {len(issues)}")
    return True
