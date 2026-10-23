"""기업 확장 정의서(extensions/{ent}/확장정의서.xlsx) 양식 — 표준 테이블 정의서와 같은 시트·열.

표준 정의서의 시트 이름과 열을 그대로 쓰고(생성기 로더 재사용), 확장 전용 열만 오른쪽에 덧붙인다.

  ext-template : 빈 양식(작성기준·예시·드롭다운 포함)
  ext-sync     : 작성자가 비워 둔 파생 행·시트를 채운다
                 (PK·기본 속성·Core 키 행, 순번, 참조속성 대상여부, 테이블목록 키·집계, 관계목록, DDL)
"""
from __future__ import annotations

import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ------------------------------------------------------------------ 시트·열(표준 열 + 확장 열)
STD_LIST = ["No", "업무영역", "조항", "표", "테이블명(국문)", "테이블명(영문)", "물리 테이블명", "테이블 설명(표준 정의)",
            "PK", "UK(식별 키)", "UK 근거", "UK 비고", "컬럼 수", "기본 속성 수", "물리 FK 수"]
EXT_LIST = ["확장유형", "대상 Core 명사", "판정 규칙(5.6)", "원천(시스템.테이블)", "관련 시나리오·KPI·질의", "상태"]
STD_TD = ["테이블명(영문)", "테이블명(국문)", "물리 테이블명", "순번", "속성명(국문)", "컬럼명(영문)", "물리 컬럼명", "구분",
          "데이터 타입(표준)", "물리 타입(권장)", "PK", "UK", "FK", "NOT NULL", "FK 참조(테이블.컬럼)", "참조속성\n대상여부",
          "코드 참조", "설명"]
EXT_TD = ["원천(시스템.테이블.컬럼)", "근거(시나리오·KPI·질의)", "상태"]
REL = ["No", "자식 테이블", "FK 컬럼", "부모 테이블", "부모 키 컬럼(UK)", "관계 구분", "비고", "FK 구분"]
GROUP = ["No", "업무영역", "공통코드그룹\n(CommonCodeGroup)", "그룹명", "코드체계", "상위 그룹", "그룹 설명",
         "사용 속성(영문)", "사용 명사", "확장 구분", "상태"]
CODE = ["No", "업무영역", "공통코드그룹\n(CommonCodeGroup)", "그룹명", "공통코드식별자\n(CommonCodeID)",
        "공통코드명\n(CommonCodeName)", "영문명", "표시순번\n(CommonCodeSequence)", "사용여부\n(UseIndicator)",
        "상위공통코드그룹\n(ParentCommonCodeGroup)", "상위공통코드식별자\n(ParentCommonCodeID)", "설명", "출처구분",
        "회사 적용\n(Y/N)", "회사 코드값", "회사 코드명", "비고(회사)", "상태"]
JUDGE = ["No", "시나리오·KPI", "원천(시스템.테이블.컬럼)", "원천 의미", "판정", "표준·확장 대상", "근거", "처리 단계", "상태"]
APPROVE = ["구분", "항목", "내용", "일자", "서명·비고"]
DDL = ["테이블명", "CREATE TABLE (PostgreSQL 참고, ext 스키마)"]

SHEETS = [("작성기준", None), ("테이블목록", STD_LIST + EXT_LIST), ("테이블정의서", STD_TD + EXT_TD),
          ("관계목록", REL), ("코드그룹목록", GROUP), ("공통코드", CODE), ("판정목록", JUDGE),
          ("검토·승인", APPROVE), ("DDL", DDL)]
WIDTHS = {
    "테이블목록": [5, 14, 7, 5, 16, 24, 24, 44, 26, 40, 14, 30, 7, 8, 8, 11, 20, 12, 28, 26, 8],
    "테이블정의서": [24, 16, 24, 6, 18, 28, 28, 8, 11, 15, 5, 5, 5, 8, 34, 9, 32, 52, 30, 28, 8],
    "관계목록": [6, 26, 40, 28, 40, 14, 50, 8],
    "코드그룹목록": [5, 14, 26, 18, 16, 20, 50, 28, 28, 16, 8],
    "공통코드": [5, 12, 26, 16, 16, 18, 20, 8, 8, 22, 16, 40, 10, 8, 12, 16, 24, 8],
    "판정목록": [5, 18, 34, 30, 12, 40, 50, 10, 8],
    "검토·승인": [12, 26, 60, 12, 30],
    "DDL": [28, 120],
}
DTYPES = ["ID", "IDREF", "enum", "string", "int", "float", "date", "dateTime", "boolean", "duration"]
PTYPE = {"ID": "VARCHAR(50)", "IDREF": "VARCHAR(50)", "enum": "VARCHAR(30)", "string": "VARCHAR(255)",
         "int": "INTEGER", "float": "NUMERIC(18,6)", "date": "DATE", "dateTime": "TIMESTAMP", "boolean": "BOOLEAN",
         "duration": "INTERVAL", "bigint": "BIGINT"}
EXT_TYPES = ["확장 명사", "확장 속성"]
KINDS = ["PK", "기본", "Core키", "확장"]
STATES = ["초안", "검토", "승인", "후보", "예시"]
JUDGEMENTS = ["표준 매핑", "코드 대응", "확장 속성", "확장 명사", "확장 코드", "제외", "S3 처리"]
GROUP_KINDS = ["신규 그룹", "표준 그룹 하위 확장"]
BASE_ROWS = ["EnterpriseID", "SiteID", "PlantID", "CreatorID", "CreationDateTime", "ModifierID",
             "ModificationDateTime", "Note"]

NAVY, WHITE, GRAY = "1F3864", "FFFFFF", "808080"
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "맑은 고딕"
EXT_HEAD = "2E75B6"     # 확장 열 머리글(표준 열과 구분)

GUIDE = [
    ("KS X 9101-1 기업 확장(Extension) 정의서", None),
    ("기준: KS X 9101-1:2025 개정(안) v1.0 테이블 정의서와 같은 양식 · 구현안내서 5.6(Core/Extension 판정), 7장(S4 생성) · "
     "이 파일이 ksxgen의 Extension 입력이다(extensions/{기업}/확장정의서.xlsx)", None),
    (None, None),
    ("작성 단계", None),
    ("S1 원천 분석", "원천 공백·Extension 후보를 찾는다(인벤토리·시나리오 정의서·KPI 정의서의 '공백·Extension')"),
    ("S2 표준 매핑", "후보마다 5.6 규칙으로 판정해 이 파일에 적고 현업 승인을 받는다(판정목록 → 테이블목록·테이블정의서·공통코드)"),
    ("S4 T-Box 생성", "ksxgen build --ext extensions/{기업} → Extension T-Box·SHACL·OBDA·DDL·코드 목록 생성(Core는 수정하지 않음)"),
    (None, None),
    ("구성", None),
    ("테이블목록", "확장 명사·확장 속성 테이블 목록. 표준 열(No~물리 FK 수) + 확장 열(확장유형, 대상 Core 명사, 판정 규칙, 원천, 관련 시나리오·KPI, 상태)"),
    ("테이블정의서", "확장 테이블의 컬럼. 표준과 같은 열. 구분 = PK / 기본 / Core키 / 확장"),
    ("관계목록", "자동 생성(ext-sync). FK 참조 열에서 도출"),
    ("코드그룹목록·공통코드", "확장 코드 그룹과 코드. 표준 공통코드 워크시트와 같은 열. '회사 코드값'은 원천 코드(대응표 겸용)"),
    ("판정목록", "S1 후보 → 5.6 판정 기록(표준 매핑·코드 대응·확장·제외·S3 처리). 확장이 아닌 판정도 근거와 함께 남긴다"),
    ("검토·승인", "ksxgen 검사 결과와 승인 서명"),
    ("DDL", "자동 생성(ext-sync). ext 스키마 CREATE TABLE 참고 스크립트"),
    (None, None),
    ("판정 규칙(5.6)", None),
    ("① 표준 매핑", "표준 속성과 의미·타입이 같다 → 확장하지 않는다(매핑 명세에 기록)"),
    ("② 확장 속성", "표준 명사에 속하지만 표준 속성에 없다 → 확장유형 '확장 속성': ext 스키마 1:1 테이블({명사}Ext), Core UK 참조"),
    ("③ 확장 명사", "표준 명사로 표현할 수 없는 업무 객체 → 확장유형 '확장 명사': ID 식별, 기본 속성 3종, 참조 대상 명시"),
    ("④ 제외", "분석·운영에 쓰지 않는 기술 컬럼 → 제외(판정목록에 사유)"),
    ("코드", "표준 그룹 코드로 대응되면 '코드 대응'(확장 아님). 회사정의(계층) 그룹에 하위 코드를 더하면 '표준 그룹 하위 확장', 표준에 없는 그룹은 '신규 그룹'"),
    (None, None),
    ("작성 방법", None),
    ("테이블목록", "확장유형·테이블명(국문·영문)·업무영역·설명·UK(확장 명사)·대상 Core 명사(확장 속성)만 적는다. "
                  "확장 속성 테이블의 영문명은 {대상 Core 명사}Ext, UK는 비워 두면 Core UK로 채운다"),
    ("테이블정의서", "업무 속성(구분 '확장')만 적어도 된다. PK·기본 속성 8종(확장 명사)·Core키(확장 속성) 행은 ext-sync가 채운다"),
    ("데이터 타입", "ID · IDREF · enum · string · int · float · date · dateTime · boolean · duration (물리 타입은 비워 두면 권장값)"),
    ("IDREF", "FK 참조(테이블.컬럼)에 부모 테이블의 UK 컬럼을 적는다(예: ItemMaster.ItemID). 부모 UK가 복합이면 컬럼마다 적는다"),
    ("enum", "코드 참조에 CommonCodeGroup=그룹(표준·확장 그룹) 또는 단위그룹(UOMGroup).UOMCode"),
    ("이름", "표준 속성과 같은 영문명은 같은 의미·타입일 때만 쓴다(그대로 공유 속성 ksx:… 재사용). 새 이름은 기업 네임스페이스로 생성"),
    ("상태", "초안 → 검토 → 승인. '후보'는 원천이 아직 없는 항목, '예시'는 생성기가 무시"),
    (None, None),
    ("금지(빌드 실패)", None),
    ("Core 수정", "표준 명사·속성·코드의 삭제·재정의, 표준 속성과 같은 이름에 다른 타입, 표준 테이블명·코드 그룹명 재사용(하위 확장 제외)"),
    ("이름 공간", "ksx: 이름 추가 금지 — 확장 이름은 기업 네임스페이스(extensions/{기업}/ext.yaml의 prefix·ns)에 만든다"),
    ("공리", "Core 클래스에 대한 동치·배타 선언, Core 속성 domain/range 재정의 금지"),
    (None, None),
    ("명령", None),
    ("양식", "python -m ksxgen ext-template --out extensions/{기업}/확장정의서.xlsx"),
    ("채움·검사", "python -m ksxgen ext-sync extensions/{기업}/확장정의서.xlsx --std standards"),
    ("생성", "python -m ksxgen build --std standards --ext extensions/{기업} --out ontologies"),
]


def _style_header(ws, ncols_std: int | None = None):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=c)
        ext = ncols_std is not None and c > ncols_std
        cell.font = Font(name=FONT, size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=EXT_HEAD if ext else NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 31.5
    ws.freeze_panes = "A2"


def _style_body(ws, start=2):
    for row in ws.iter_rows(min_row=start, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _widths(ws, name):
    for i, w in enumerate(WIDTHS.get(name, []), 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _dv(ws, col_name, header, values, rows=2000):
    if col_name not in header:
        return
    col = get_column_letter(header.index(col_name) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "목록에서 고르세요", "입력값 확인"
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{rows}")


def _guide_sheet(wb):
    ws = wb.create_sheet("작성기준", 0)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 120
    for i, (a, b) in enumerate(GUIDE, 2):
        if a is not None:
            ws.cell(row=i, column=2, value=a)
        if b is not None:
            ws.cell(row=i, column=3, value=b)
        c2 = ws.cell(row=i, column=2)
        if b is None and a:
            c2.font = Font(name=FONT, size=14 if i == 2 else 11, bold=True, color=NAVY)
        else:
            c2.font = Font(name=FONT, size=10, bold=True)
            ws.cell(row=i, column=3).font = Font(name=FONT, size=10)
            ws.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def new_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _guide_sheet(wb)
    for name, header in SHEETS[1:]:
        ws = wb.create_sheet(name)
        ws.append(header)
        nstd = {"테이블목록": len(STD_LIST), "테이블정의서": len(STD_TD), "공통코드": len(CODE) - 1,
                "코드그룹목록": len(GROUP) - 2}.get(name)
        _style_header(ws, nstd)
        _widths(ws, name)
    return wb


def add_validations(wb):
    ws = wb["테이블목록"]
    _dv(ws, "확장유형", STD_LIST + EXT_LIST, EXT_TYPES)
    _dv(ws, "판정 규칙(5.6)", STD_LIST + EXT_LIST, ["② 확장 속성", "③ 확장 명사"])
    _dv(ws, "상태", STD_LIST + EXT_LIST, STATES)
    ws = wb["테이블정의서"]
    h = STD_TD + EXT_TD
    _dv(ws, "구분", h, KINDS)
    _dv(ws, "데이터 타입(표준)", h, DTYPES)
    for c, vals in (("PK", ["Y"]), ("UK", ["UK"]), ("FK", ["Y"]), ("NOT NULL", ["Y"]),
                    ("참조속성\n대상여부", ["대상", "비대상"]), ("상태", STATES)):
        _dv(ws, c, h, vals)
    _dv(wb["코드그룹목록"], "확장 구분", GROUP, GROUP_KINDS)
    _dv(wb["코드그룹목록"], "상태", GROUP, STATES)
    _dv(wb["공통코드"], "사용여부\n(UseIndicator)", CODE, ["Y", "N"])
    _dv(wb["공통코드"], "회사 적용\n(Y/N)", CODE, ["Y", "N"])
    _dv(wb["공통코드"], "상태", CODE, STATES)
    _dv(wb["판정목록"], "판정", JUDGE, JUDGEMENTS)
    _dv(wb["판정목록"], "처리 단계", JUDGE, ["S2", "S3", "S4", "원천 개선"])
    _dv(wb["판정목록"], "상태", JUDGE, STATES)


EXAMPLE_LIST = [None, "품질관리(확장)", "확장", None, "예시명사", "ExampleNoun", "ExampleNoun",
                "(예시) 표준 명사로 표현할 수 없는 업무 객체의 정의를 적는다", None, "EnterpriseID + ExampleNounID",
                "확장 ID", None, None, None, None, "확장 명사", None, "③ 확장 명사", "erp.tb_example",
                "SC00 / KPI00 / Q-SC00-01", "예시"]
EXAMPLE_TD = [["ExampleNoun", "예시명사", "ExampleNoun", None, "예시명사식별자", "ExampleNounID", "ExampleNounID", "확장",
               "ID", "VARCHAR(50)", None, "UK", None, "Y", None, None, None, "— (예시) 확장 명사의 식별자",
               "erp.tb_example.ex_cd", "KPI00", "예시"],
              ["ExampleNoun", "예시명사", "ExampleNoun", None, "품목식별자", "ItemID", "ItemID", "확장", "IDREF",
               "VARCHAR(50)", None, None, "Y", "Y", "ItemMaster.ItemID", "대상", None,
               "— (예시) 표준 속성 이름을 같은 의미로 재사용\n— 참조 대상: 품목마스터(ItemMaster)",
               "erp.tb_example.item_cd", "KPI00", "예시"]]


def template(path: str):
    wb = new_workbook()
    gray = Font(name=FONT, size=10, italic=True, color=GRAY)
    ws = wb["테이블목록"]
    ws.append(EXAMPLE_LIST)
    ws2 = wb["테이블정의서"]
    for r in EXAMPLE_TD:
        ws2.append(r)
    wb["코드그룹목록"].append([None, "품질관리(확장)", "ExampleType", "예시유형", "회사정의", None,
                           "(예시) 표준에 없는 코드 그룹", "ExampleType", "ExampleNoun", "신규 그룹", "예시"])
    wb["공통코드"].append([None, "품질관리(확장)", "ExampleType", "예시유형", "EX1", "예시코드", "example", 1, "Y", None, None,
                       "(예시)", "확장", "Y", "X1", "원천 코드명", None, "예시"])
    wb["판정목록"].append([None, "SC00 / KPI00", "erp.tb_example.col", "(예시) 원천 의미", "확장 속성", "ItemMaster → ItemMasterExt.ExampleAttr",
                       "② 표준 명사에 속하나 표준 속성 없음", "S2", "예시"])
    for name in ("테이블목록", "테이블정의서", "코드그룹목록", "공통코드", "판정목록"):
        for row in wb[name].iter_rows(min_row=2):
            for c in row:
                c.font = gray
                c.border = BORDER
                c.alignment = Alignment(vertical="top", wrap_text=True)
    ap = wb["검토·승인"]
    for r in [("검사", "ksxgen 검사 결과", "(ext-sync 실행 시 자동 기록)", None, None),
              ("승인", "도메인(PI) 컨설턴트", None, None, None), ("승인", "고객사 현업 검수자", None, None, None),
              ("승인", "온톨로지 엔지니어", None, None, None), ("승인", "PM", None, None, None)]:
        ap.append(list(r))
    _style_body(ap)
    add_validations(wb)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path
