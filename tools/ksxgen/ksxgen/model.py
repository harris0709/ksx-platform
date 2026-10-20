"""KS X 9101 테이블 정의서·공통코드 워크시트를 읽어 생성기 내부 모델을 만든다.

입력은 standards/ 폴더의 두 엑셀 파일이며, 외부 라이브러리는 openpyxl만 쓴다.
"""
from __future__ import annotations

import glob
import os
import re
from dataclasses import dataclass, field

import openpyxl

BASE_COLS = ("EnterpriseID", "SiteID", "PlantID")
COPY_RE = re.compile(r"참조 대상:\s*[^()\n]*\(([A-Za-z0-9]+)\.([A-Za-z0-9]+)\)")
TARGET_RE = re.compile(r"참조 대상:\s*([^()\n]*)\(([A-Za-z0-9]+)\)")


@dataclass
class Column:
    table: str
    seq: int
    ko: str
    en: str
    kind: str            # PK / 기본 / 명사
    dtype: str           # 표준 데이터 타입
    ptype: str           # 물리 타입(권장)
    pk: bool
    uk: bool
    fk: bool
    notnull: bool
    fkref: str | None
    ref: str | None      # 참조속성 대상여부
    code: str | None     # 코드 참조(CommonCodeGroup=...)
    desc: str

    @property
    def code_group(self) -> str | None:
        if self.code and "=" in self.code:
            return self.code.split("=", 1)[1].strip()
        return None

    @property
    def copy_of(self) -> tuple[str, str] | None:
        m = COPY_RE.search(self.desc or "")
        return (m.group(1), m.group(2)) if m else None


@dataclass
class Relation:
    no: int
    child: str
    fk_cols: list[str]
    parent: str
    parent_cols: list[str]
    kind: str            # 기본 속성 / 명사 참조 / 헤더-라인 / 사본 참조 ...
    note: str | None
    physical: bool


@dataclass
class Table:
    en: str
    ko: str
    area: str
    clause: str
    table_no: int | None
    definition: str
    uk: list[str]
    columns: list[Column] = field(default_factory=list)
    # 헤더·라인 구조체
    role: str = "noun"           # noun / header / line
    noun: str = ""               # 소속 명사(헤더·라인은 원 명사명)
    header: str | None = None    # 라인일 때 헤더 테이블명

    @property
    def cls(self) -> str:
        """온톨로지 클래스 로컬명: 헤더 → 명사, 라인 → ~Line, 그 외 → 테이블명."""
        if self.role == "header":
            return self.noun
        return self.en

    def col(self, en: str) -> Column | None:
        return next((c for c in self.columns if c.en == en), None)


@dataclass
class Code:
    group: str
    code: str
    name: str
    en: str | None
    seq: int | None
    use: str | None
    parent_group: str | None
    parent_code: str | None
    desc: str | None


@dataclass
class CodeGroup:
    group: str
    name: str
    area: str
    scheme: str | None
    parent: str | None
    desc: str | None


@dataclass
class Uom:
    group: str
    group_name: str
    code: str
    name: str
    en: str | None
    cefact: str | None
    base: bool
    factor: float | None


@dataclass
class Model:
    tables: dict[str, Table]
    relations: list[Relation]
    code_groups: dict[str, CodeGroup]
    codes: list[Code]
    uoms: list[Uom]
    source_files: dict[str, str]

    def noun_tables(self) -> list[Table]:
        return list(self.tables.values())

    def classes(self) -> list[Table]:
        """클래스를 만드는 테이블(헤더는 명사 클래스, 라인은 라인 클래스)."""
        return list(self.tables.values())

    @property
    def noun_count(self) -> int:
        return sum(1 for t in self.tables.values() if t.role != "line")


def _find(std_dir: str, pattern: str) -> str:
    hits = sorted(glob.glob(os.path.join(std_dir, pattern)))
    hits = [h for h in hits if not os.path.basename(h).startswith("~$")]
    if not hits:
        raise FileNotFoundError(f"{std_dir}에서 {pattern} 파일을 찾지 못했습니다")
    return hits[-1]


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [(_s(h).replace("\n", "")) for h in next(it)]
    for r in it:
        if r is None or all(v is None for v in r):
            continue
        yield dict(zip(head, r))


def load(std_dir: str) -> Model:
    td = _find(std_dir, "*테이블정의서*.xlsx")
    cw = _find(std_dir, "*공통코드*.xlsx")
    wb = openpyxl.load_workbook(td, read_only=True, data_only=True)

    tables: dict[str, Table] = {}
    for r in _rows(wb["테이블목록"]):
        en = _s(r["테이블명(영문)"])
        if not en:
            continue
        uk = [x.strip() for x in _s(r["UK(식별 키)"]).split("+") if x.strip()]
        tno = r.get("표")
        tables[en] = Table(en=en, ko=_s(r["테이블명(국문)"]), area=_s(r["업무영역"]), clause=_s(r["조항"]),
                           table_no=int(tno) if isinstance(tno, (int, float)) else None,
                           definition=_s(r["테이블 설명(표준 정의)"]), uk=uk)

    for r in _rows(wb["테이블정의서"]):
        t = tables.get(_s(r["테이블명(영문)"]))
        if t is None:
            continue
        t.columns.append(Column(
            table=t.en, seq=int(r["순번"]), ko=_s(r["속성명(국문)"]), en=_s(r["컬럼명(영문)"]), kind=_s(r["구분"]),
            dtype=_s(r["데이터 타입(표준)"]), ptype=_s(r["물리 타입(권장)"]), pk=_s(r["PK"]) == "Y",
            uk=_s(r["UK"]) == "UK", fk=_s(r["FK"]) == "Y", notnull=_s(r["NOT NULL"]) == "Y",
            fkref=_s(r["FK 참조(테이블.컬럼)"]) or None, ref=_s(r["참조속성대상여부"]) or None,
            code=_s(r["코드 참조"]) or None, desc=_s(r["설명"])))
    for t in tables.values():
        t.columns.sort(key=lambda c: c.seq)

    rels: list[Relation] = []
    for r in _rows(wb["관계목록"]):
        if not _s(r["자식 테이블"]):
            continue
        rels.append(Relation(
            no=int(r["No"]), child=_s(r["자식 테이블"]),
            fk_cols=[x.strip() for x in _s(r["FK 컬럼"]).split(",") if x.strip()],
            parent=_s(r["부모 테이블"]),
            parent_cols=[x.strip() for x in _s(r["부모 키 컬럼(UK)"]).split(",") if x.strip()],
            kind=_s(r["관계 구분"]), note=_s(r["비고"]) or None, physical=_s(r["FK 구분"]) == "물리"))

    # 헤더·라인 구조체 판정
    for t in tables.values():
        t.noun = t.en
    for rel in rels:
        if rel.kind == "헤더-라인":
            line, head = tables[rel.child], tables[rel.parent]
            noun = head.en[:-len("Header")] if head.en.endswith("Header") else head.en
            head.role, head.noun = "header", noun
            line.role, line.noun, line.header = "line", noun, head.en

    # 공통코드 워크시트
    wc = openpyxl.load_workbook(cw, read_only=True, data_only=True)
    groups: dict[str, CodeGroup] = {}
    for r in _rows(wc["코드그룹목록"]):
        g = _s(r.get("공통코드그룹(CommonCodeGroup)"))
        if g:
            groups[g] = CodeGroup(g, _s(r.get("그룹명")), _s(r.get("업무영역")), _s(r.get("코드체계")) or None,
                                  _s(r.get("상위 그룹")) or None, _s(r.get("그룹 설명")) or None)
    codes: list[Code] = []
    for r in _rows(wc["공통코드"]):
        g, c = _s(r.get("공통코드그룹(CommonCodeGroup)")), _s(r.get("공통코드식별자(CommonCodeID)"))
        if g and c:
            seq = r.get("표시순번(CommonCodeSequence)")
            codes.append(Code(g, c, _s(r.get("공통코드명(CommonCodeName)")), _s(r.get("영문명")) or None,
                              int(seq) if isinstance(seq, (int, float)) else None, _s(r.get("사용여부(UseIndicator)")) or None,
                              _s(r.get("상위공통코드그룹(ParentCommonCodeGroup)")) or None,
                              _s(r.get("상위공통코드식별자(ParentCommonCodeID)")) or None, _s(r.get("설명")) or None))
    uoms: list[Uom] = []
    for r in _rows(wc["단위코드"]):
        code = _s(r.get("단위코드(UOMCode)"))
        if code:
            f = r.get("기준단위 환산계수")
            uoms.append(Uom(_s(r.get("단위그룹식별자(UOMGroupID)")), _s(r.get("단위그룹명(UOMGroupName)")), code,
                            _s(r.get("단위코드명(UOMCodeName)")), _s(r.get("영문명")) or None,
                            _s(r.get("UN/CEFACTRec 20")) or None, _s(r.get("그룹 기준단위")) == "Y",
                            float(f) if isinstance(f, (int, float)) else None))
    return Model(tables, rels, groups, codes, uoms, {"table_definition": td, "code_worksheet": cw})
