"""S2 산출물: 매핑 명세(YAML), 코드·단위·식별 대응표(CSV), 매핑정의서(xlsx 검토본), 요약(md)."""
from __future__ import annotations

import csv
import os
import sys
from datetime import datetime

import yaml

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_HERE, "..", "..", "s1profiler"))
from s1profiler.output import _sheet  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402

STAGE_KO = {"exact": "① 정확", "normalized": "② 정규화", "component": "③ 컴포넌트", "korean": "④ 국문명",
            "manual": "⑥⑦ 검토자", "const": "상수", "xref": "식별 대응", "derived": "파생"}
FILL = {"미판정": "F8D7DA", "제외": "EEF2F7", "조인 키": "EAF6EC", "확장": "E8F1FB"}


class _Flow(dict):
    pass


def _flow_rep(dumper, data):
    return dumper.represent_mapping("tag:yaml.org,2002:map", data, flow_style=True)


yaml.add_representer(_Flow, _flow_rep, Dumper=yaml.SafeDumper)


def _spec_yaml(eng, r) -> str:
    sp, noun = r.spec, eng.std.nouns[r.noun]
    src = {"system": r.system, "table": r.table}
    if sp.get("filter"):
        src["filter"] = sp["filter"]
    doc = {"noun": r.noun, "layer": noun.layer, "structure": sp.get("structure", "1:1"), "source": _Flow(src)}
    if sp.get("join"):
        doc["join"] = [_Flow({k: v for k, v in j.items() if k != "exclude"}) for j in sp["join"]]
    doc["base"] = _Flow({a: m.transform for a, m in r.maps.items() if a in ("EnterpriseID", "SiteID", "PlantID")})
    doc["identity"] = [k for k in noun.uk if k != "EnterpriseID"]
    attrs, ext = {}, {}
    for a, m in r.maps.items():
        if a in ("EnterpriseID", "SiteID", "PlantID"):
            continue
        d = {}
        if m.src:
            d["from"] = m.src
        d["match"] = m.stage
        for k in ("code", "uom", "transform", "note"):
            v = getattr(m, k)
            if v:
                d[k if k != "code" else "code"] = (f"{v} -> {noun.attrs[a].code_group}" if k == "code" else v)
        d["status"] = m.status
        (ext if m.layer == "ext" else attrs)[a] = _Flow(d)
    doc["attributes"] = attrs
    doc["extension"] = ext
    doc["excluded"] = [_Flow({"column": c, "reason": w}) for c, w in r.excluded.items()]
    if r.unresolved:
        doc["unresolved"] = [_Flow({"column": c, "candidates": [a.en for a, _ in cs[:3]]}) for c, cs in r.unresolved.items()]
    if sp.get("note"):
        doc["note"] = sp["note"]
    head = (f"# {r.noun} ← {r.table} ({doc['structure']}) — automap 생성, 계획: s2_plan.yaml\n"
            f"# 상태: draft(자동) → reviewed(데이터 엔지니어) → approved(도메인·현업)\n")
    return head + yaml.dump(doc, Dumper=yaml.SafeDumper, allow_unicode=True, sort_keys=False, width=160)


def spec_path(out, r):
    sp = r.spec
    name = sp.get("file") or r.noun
    return os.path.join(out, r.system.lower(), f"{name}.yaml")


def write_all(eng, out: str) -> list[str]:
    ent = eng.plan["enterprise"]
    paths = []
    for r in eng.results:
        p = spec_path(out, r)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        open(p, "w", encoding="utf-8").write(_spec_yaml(eng, r))
        paths.append(p)
    # 대응표 CSV
    rows = eng.code_rows()
    with open(os.path.join(out, "codes.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["src_system", "src_group", "src_code", "ksx_group", "ksx_code", "rule", "note"])
        for x in rows:
            w.writerow([x["src_system"], x["src_group"], x["src_code"], x["ksx_group"], x["ksx_code"], x["rule"], x["note"]])
    with open(os.path.join(out, "units.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["src_unit", "ksx_uom", "factor", "note"])
        for u in eng.plan.get("units", []):
            w.writerow([u["src_unit"], u["ksx_uom"], u.get("factor", 1), u.get("note", "")])
    with open(os.path.join(out, "xref.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["xref", "src_system", "src", "ksx", "note"])
        for name, items in eng.plan.get("xref", {}).items():
            for x in items:
                w.writerow([name, x.get("system", ""), x["src"], x["ksx"], x.get("note", "")])
    paths += [os.path.join(out, f) for f in ("codes.csv", "units.csv", "xref.csv")]
    # 검토본
    judg = eng.judgements()
    met = eng.metrics(judg)
    paths.append(_xlsx(eng, out, ent, judg, met))
    paths.append(_md(eng, out, ent, judg, met))
    return paths


def _xlsx(eng, out, ent, judg, met) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    summ = [("대상 기업", ent), ("생성", stamp + " · tools/automap · 계획 mapping/" + ent.lower() + "/s2_plan.yaml"),
            ("매핑 명세(원천 → 명사)", met["specs"]), ("대상 명사(Core+확장)", met["nouns"]),
            ("매핑 속성(원천 컬럼 1개 기반)", met["mapped"]),
            ("파생 속성(여러 컬럼 조합 — 식별자 생성·계산)", met["derived"]),
            ("결정론 해소(①~④)", f"{met['deterministic']} ({met['det_rate']}%) — 목표 70% 이상"),
            ("검토자 수정(결정론 후보와 다르게 확정)", f"{met['corrected']} ({met['fix_rate']}%) — 목표 20% 미만"),
            ("검토자 매핑(결정론 미해소 → ⑤~⑦)", met["manual"]),
            ("필수 속성(UK·NOT NULL·기본 3종) 매핑", f"{met['req_total'] - met['req_missing']}/{met['req_total']} ({met['req_rate']}%) — 완료 기준 100%"),
            ("대상 컬럼 판정", f"{met['judged']}/{met['cols']} ({met['judge_rate']}%) — 완료 기준 100%"),
            ("판정 분포", ", ".join(f"{k} {v}" for k, v in sorted(met["by_kind"].items()))),
            ("검사 오류", met["errors"]),
            ("상태", "reviewed(데이터 엔지니어 검토) — 도메인·현업 승인 대기")]
    _sheet(wb, "요약", ["항목", "값"], summ, [34, 90], wrap_cols=(2,))
    _sheet(wb, "구조매핑", ["No", "명사", "구분", "구조 유형", "원천 시스템", "원천 테이블", "조인", "필터", "기본 속성", "식별(UK)",
                         "매핑 속성", "필수 누락", "명세 파일", "비고"],
           [(i, r.noun, eng.std.nouns[r.noun].layer, r.spec.get("structure", "1:1"), r.system, r.table,
             "; ".join(f"{j['table']} ON {j['on']}" for j in r.spec.get("join", [])), r.spec.get("filter", ""),
             ", ".join(f"{a}={r.maps[a].transform}" for a in ("EnterpriseID", "SiteID", "PlantID") if a in r.maps),
             " + ".join(eng.std.nouns[r.noun].uk), len(r.maps), ", ".join(r.missing),
             os.path.relpath(spec_path(out, r), eng.repo), r.spec.get("note", ""))
            for i, r in enumerate(eng.results, 1)],
           [5, 24, 6, 12, 7, 20, 34, 18, 36, 34, 7, 16, 30, 40], wrap_cols=(7, 9, 10, 14))
    rows = []
    for r in eng.results:
        noun = eng.std.nouns[r.noun]
        for a in noun.attrs.values():
            m = r.maps.get(a.en)
            if m is None and not a.required:
                continue
            tb, col = r.resolve(m.src) if (m and m.src and "," not in m.src and "+" not in m.src) else (r.table, None)
            rows.append((r.noun, a.en, a.ko, a.dtype, "UK" if a.uk else ("필수" if a.required else ""),
                         "확장" if a.layer == "ext" else "Core", (m.src if m else ""), STAGE_KO.get(m.stage, m.stage) if m else "누락",
                         (m.evidence or "") if m else "", ((m.code + " → " + (a.code_group or "")) if m and m.code else ""),
                         (m.uom or "") if m else "", (m.transform or "") if m else "", eng.samples(tb, col) if col else "",
                         (m.status if m else "미매핑"), (m.note or "") if m else ""))
    ws = _sheet(wb, "속성매핑", ["명사", "표준 속성", "국문", "타입", "키·필수", "구분", "원천 컬럼", "매칭 단계", "매칭 근거",
                              "코드 대응", "단위", "변환", "원천 표본 값(상위 5)", "상태", "비고"], rows,
                [22, 26, 16, 8, 7, 6, 22, 11, 34, 26, 14, 30, 36, 9, 34], wrap_cols=(9, 12, 13, 15))
    for row in ws.iter_rows(min_row=2):
        if row[7].value == "누락":
            for c in row:
                c.fill = PatternFill("solid", fgColor="F8D7DA")
    ws = _sheet(wb, "컬럼판정", ["시스템.테이블", "컬럼", "주석", "판정", "대상(명사.속성)", "사유", "표본 값"],
                [(j["table"], j["column"], j["comment"], j["kind"], j["targets"], j["reason"], j["samples"]) for j in judg],
                [22, 16, 22, 9, 46, 46, 36], wrap_cols=(5, 6, 7))
    for row in ws.iter_rows(min_row=2):
        f = FILL.get(row[3].value)
        if f:
            row[3].fill = PatternFill("solid", fgColor=f)
    _sheet(wb, "후보(미해소)", ["명사", "원천 컬럼", "후보 1", "후보 2", "후보 3", "후보 4", "후보 5"],
           [(r.noun, c) + tuple(f"{a.en}({s})" for a, s in cs) for r in eng.results for c, cs in r.unresolved.items()],
           [22, 20, 26, 26, 26, 26, 26])
    _sheet(wb, "코드대응", ["원천 시스템", "원천 그룹", "원천 코드", "표준 그룹", "표준 코드", "표준 코드명", "규칙", "비고"],
           [(x["src_system"], x["src_group"], x["src_code"], x["ksx_group"], x["ksx_code"],
             eng.std.codes.get(x["ksx_group"], {}).get(x["ksx_code"], "?"), x["rule"], x["note"]) for x in eng.code_rows()],
           [8, 14, 12, 26, 14, 18, 14, 40], wrap_cols=(8,))
    _sheet(wb, "단위대응", ["원천 단위", "표준 단위코드", "단위명", "환산 계수", "비고"],
           [(u["src_unit"], u["ksx_uom"], eng.std.uoms.get(u["ksx_uom"], "?"), u.get("factor", 1), u.get("note", ""))
            for u in eng.plan.get("units", [])], [12, 14, 16, 10, 40])
    _sheet(wb, "식별대응", ["대응표", "원천 시스템", "원천 값", "표준 식별자", "비고"],
           [(n, x.get("system", ""), x["src"], x["ksx"], x.get("note", "")) for n, items in eng.plan.get("xref", {}).items()
            for x in items], [14, 10, 16, 20, 50])
    _sheet(wb, "테이블판정", ["시스템.테이블", "판정", "비고"],
           [(t, d.get("disposition", ""), d.get("note", "")) for t, d in eng.plan.get("tables", {}).items()], [26, 14, 80],
           wrap_cols=(3,))
    _sheet(wb, "검사", ["수준", "위치", "내용"], eng.issues or [("—", "—", "오류·경고 없음")], [8, 40, 90], wrap_cols=(3,))
    _sheet(wb, "승인", ["역할", "이름", "일자", "의견"],
           [("데이터 엔지니어(reviewed)", "", "", ""), ("도메인(PI) 컨설턴트(approved)", "", "", ""),
            ("고객사 현업 검수자(approved)", "", "", "")], [30, 16, 12, 50])
    p = os.path.join(out, f"S2_매핑정의서_{ent}.xlsx")
    for ws in wb:                                   # 인쇄: 가로·폭 맞춤·제목 행 반복
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_title_rows = "1:1"
    wb.save(p)
    return p


def _md(eng, out, ent, judg, met) -> str:
    md = [f"# S2 표준 매핑 요약 — {ent}", "", "| 항목 | 값 |", "|---|---|",
          f"| 매핑 명세 / 대상 명사 | {met['specs']} / {met['nouns']} |",
          f"| 매핑 속성(컬럼 1개) / 파생(여러 컬럼) | {met['mapped']} / {met['derived']} |",
          f"| 결정론 해소율 | {met['det_rate']}% (목표 70% 이상) |",
          f"| 검토자 수정 비율 | {met['fix_rate']}% (목표 20% 미만) |",
          f"| 필수 속성 매핑 | {met['req_rate']}% ({met['req_total'] - met['req_missing']}/{met['req_total']}) |",
          f"| 대상 컬럼 판정 | {met['judge_rate']}% ({met['judged']}/{met['cols']}) |",
          f"| 판정 분포 | {', '.join(f'{k} {v}' for k, v in sorted(met['by_kind'].items()))} |",
          f"| 검사 오류 | {met['errors']} |", "", "## 명세", "", "| 명사 | 원천 | 구조 | 매핑 | 필수 누락 |", "|---|---|---|---|---|"]
    md += [f"| {r.noun} | {r.table} | {r.spec.get('structure', '1:1')} | {len(r.maps)} | {', '.join(r.missing) or '—'} |"
           for r in eng.results]
    md += ["", "## 검사", ""] + ([f"- {a} · {b}: {c}" for a, b, c in eng.issues] or ["- 오류 없음"])
    p = os.path.join(out, f"S2_매핑요약_{ent}.md")
    open(p, "w", encoding="utf-8").write("\n".join(md) + "\n")
    return p
