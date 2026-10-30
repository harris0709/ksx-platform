"""정제 규칙 검토서(구현안내서 6.7) — 현업과 함께 처리 방식을 정하는 회의 자료.

입력: s3_plan.yaml의 rules·review·s2_regressions·approval + 마지막 publish(커밋)의 audit 기록
출력: mapping/<기업>/S3_정제규칙_검토서_<기업>.xlsx

  시트          내용
  작성기준      검토 절차·처리 방식·원칙
  검토목록      규칙 1행: 이슈·대상·처리 행 수·근거·원천 확인·선택지·결정·원천 정정 요청
  처리행        규칙별 처리 전 원래 행(audit.quarantine) — 현업이 원천과 대조
  대사          모델별 행 수·합계 대사(audit.load_log)
  원천정정요청  데이터 소유 부서에 보낼 요청(원천 시스템·테이블 포함)
  S2회귀        S3에서 S2 명세로 되돌아간 결정
  승인          결정·이관 과제·서명

판단(원천 확인·선택지·결정)은 계획 파일에 적고, 이 검토서는 언제든 다시 만든다(S2 검토본과 같은 원칙).
"""
from __future__ import annotations

import os
import re

import yaml

from .base import Config, rows

ACTION_KO = {
    "dedupe": "중복 정리", "quarantine": "격리", "nullify_orphan": "참조만 비움",
    "replace_orphan": "다른 값으로 대체", "correct": "보정", "flag": "표시(현업 확인)",
}
ACTION_DESC = [
    ("dedupe", "같은 대상이 표기만 달리 여러 번 등록", "1행만 남김"),
    ("quarantine", "값을 믿을 수 없고 바로잡을 근거도 없음", "행을 적재하지 않음"),
    ("nullify_orphan", "행은 맞지만 참조 대상이 없음", "해당 참조 컬럼만 NULL"),
    ("replace_orphan", "같은 업무의 다른 기록에서 올바른 참조를 찾을 수 있음", "참조 컬럼을 근거 값으로 바꿈"),
    ("correct", "서로 독립된 근거 2개 이상이 같은 값을 가리킬 때만", "값을 고침(보정 컬럼은 합계 대사에서 뺌)"),
    ("flag", "판단을 미루고 현업 확인만 요청", "그대로 적재"),
]
BASE_COLS = {"EnterpriseID", "SiteID"}
APPROVAL_KO = {"pending": "대기", "approved": "승인", "conditional": "조건부 승인", "rejected": "반려"}


def _parent_review(rid: str, review: dict):
    """R07b → R07처럼 같은 이슈의 파생 규칙은 앞 규칙의 검토를 따른다."""
    if rid in review:
        return review[rid], rid
    base = re.sub(r"[a-z]+$", "", rid)
    if base != rid and base in review:
        return review[base], base
    return {}, None


def _source_of(cfg: Config, model: str, manifest: dict) -> str:
    m = next((x for x in manifest["models"] if x["model"] == model), None)
    if not m:
        return ""
    spec = yaml.safe_load(open(os.path.join(cfg.repo, m["spec"]), encoding="utf-8"))
    src = spec.get("source", {})
    return f"{src.get('system', '')} {src.get('table', '')}".strip()


def build_review(cfg: Config, out: str | None = None, run_id: str | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    raw = cfg.raw
    rules = raw.get("rules", [])
    review = raw.get("review", {}) or {}
    regress = raw.get("s2_regressions", []) or []
    appr = raw.get("approval", {}) or {}
    manifest = yaml.safe_load(open(os.path.join(cfg.dbt_dir, "publish", f"{cfg.ent.lower()}.yml"), encoding="utf-8"))

    if not run_id:
        r = rows(cfg.owner, "select run_id::text rid, started_at::text t from audit.run where stage = 'S3' and job = 'publish' "
                            f"and status = 'succeeded' and enterprise_id = '{cfg.ent}' "
                            "and coalesce(detail->>'committed', 'false') = 'true' order by started_at desc limit 1")
        if not r:
            raise SystemExit("ksxload review: 커밋된 publish 기록이 없음 — publish를 먼저 실행")
        run_id, run_at = r[0]["rid"], r[0]["t"][:19]
    else:
        run_at = rows(cfg.owner, f"select started_at::text t from audit.run where run_id = '{run_id}'")[0]["t"][:19]
    q = rows(cfg.owner, f"select rule_id, action, model, row_data from audit.quarantine where run_id = '{run_id}' order by q_id")
    logs = rows(cfg.owner, f"select object, source_rows, target_rows, status, detail from audit.load_log "
                           f"where run_id = '{run_id}' and step = 'check' order by object")
    cnt = {}
    for x in q:
        cnt[x["rule_id"]] = cnt.get(x["rule_id"], 0) + 1

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="DCE6F2")
    input_fill = PatternFill("solid", fgColor="FFF7E0")      # 현업·검토자가 판단하는 칸
    thin = Side(style="thin", color="A0A8B5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(title, header, data, widths, inputs=()):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = input_fill if c.column - 1 in inputs else head_fill
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for row in data:
            ws.append(row)
        for r_ in ws.iter_rows(min_row=2):
            for c in r_:
                c.border = border
                c.alignment = wrap
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    # ---- 작성기준
    ws = wb.active
    ws.title = "작성기준"
    guide = [
        [f"S3 정제 규칙 검토서 — {cfg.ent}"],
        [f"기준 적재: run {run_id} ({run_at}) · 규칙 {len(rules)}개 · 처리 행 {len(q)}"],
        [],
        ["절차(구현안내서 6.7)"],
        ["1. 첫 적재(publish)가 대사를 통과한 뒤, S3 승인 전에 검토한다"],
        ["2. 규칙을 하나씩 순서대로: ① 원천 데이터(처리행 시트) ② 교차 확인 ③ 분석 ④ 선택지 ⑤ 재현 쿼리"],
        ["3. 현업이 사실(입력 오류·공용 계정 등)을 확인하고 처리 방식을 정한다 — 노란 머리글 열"],
        ["4. 결정은 s3_plan.yaml의 review 절에 적고 규칙을 고친 뒤 gen → dbt build → publish → review로 다시 만든다"],
        ["5. 원천 정정 요청을 데이터 소유 부서에 보내고, 승인 시트에 서명한다"],
        [],
        ["처리 방식", "언제 쓰나", "원장에 미치는 영향"],
        *[[f"{ACTION_KO[a]} ({a})", w, e] for a, w, e in ACTION_DESC],
        [],
        ["원칙"],
        ["· 모든 처리는 원래 행을 audit.quarantine에 남긴다. 원천은 바꾸지 않는다"],
        ["· 보정·대체는 그럴듯한 값이 아니라 원천 안의 다른 기록에서 가져오고, 근거를 규칙 note에 적는다"],
        ["· 같은 원인이 여러 테이블에 퍼져 있으면 한 번에 같은 방식으로 결정한다"],
        ["· 원장 처리와 별도로 원천 정정 요청을 함께 정한다"],
    ]
    for g in guide:
        ws.append(g)
    ws["A1"].font = Font(bold=True, size=13)
    for r_ in (4, 11, 11 + len(ACTION_DESC) + 2):
        for c in ws[r_]:
            c.font = Font(bold=True)
    for c in ws[11]:
        c.fill = head_fill
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 40

    # ---- 검토목록
    data = []
    for r_ in rules:
        rv, of = _parent_review(r_["id"], review)
        follow = f"{of}과 같음 — " if of and of != r_["id"] else ""
        data.append([
            r_["id"], r_.get("issue", ""), r_["model"], ACTION_KO.get(r_["action"], r_["action"]), cnt.get(r_["id"], 0),
            r_.get("note", ""), rv.get("finding", ""), rv.get("options", ""),
            (follow + rv["decision"]) if rv.get("decision") else "", rv.get("decided", ""),
            "결정" if rv.get("decision") else "미결정",
        ])
    sheet("검토목록", ["규칙", "이슈", "대상 모델", "처리 방식", "처리 행", "근거(규칙 note)", "원천 확인 결과",
                    "선택지", "결정", "결정일", "상태"],
          data, [7, 8, 22, 14, 8, 40, 44, 36, 40, 11, 8], inputs=(6, 7, 8, 9))

    # ---- 처리행: 규칙별로 컬럼 머리글을 따로
    ws = wb.create_sheet("처리행")
    ws.append([f"처리 전 원래 행(audit.quarantine, run {run_id}) — 원천과 대조해 사실을 확인한다"])
    ws["A1"].font = Font(bold=True)
    widths = {}
    for r_ in rules:
        items = [x for x in q if x["rule_id"] == r_["id"]]
        ws.append([])
        ws.append([f"{r_['id']} · {r_.get('issue', '')} · {r_['model']} · {ACTION_KO.get(r_['action'], r_['action'])} · {len(items)}행"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        if not items:
            ws.append(["(처리 행 없음)"])
            continue
        keys = list(items[0]["row_data"])                 # jsonb는 키 순서를 바꾸므로 원장 컬럼 순서로
        mm = next((m for m in manifest["models"] if m["model"] == r_["model"]), None)
        order = list((mm or {}).get("uk", [])) + [c for t in (mm or {}).get("targets", []) for c in t["columns"]]
        cols = [k for k in order if k in keys] + [k for k in keys if k not in order]
        cols = list(dict.fromkeys(k for k in cols if k not in BASE_COLS))
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.fill = head_fill
            c.font = Font(bold=True)
            c.border = border
        for x in items:
            ws.append([x["row_data"].get(k) if not isinstance(x["row_data"].get(k), (dict, list)) else str(x["row_data"].get(k))
                       for k in cols])
            for c in ws[ws.max_row]:
                c.border = border
        for i, k in enumerate(cols, 1):
            widths[i] = max(widths.get(i, 10), min(28, len(k) + 2))
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---- 대사
    data = []
    for lg in logs:
        d = lg["detail"] or {}
        sums = d.get("sums", [])
        rules_ = ", ".join(f"{k} {v}" for k, v in (d.get("rules") or {}).items())
        data.append([lg["object"], lg["source_rows"], d.get("removed_by_rules", 0), lg["target_rows"],
                     f"{sum(1 for s in sums if s.get('ok'))}/{len(sums)}" if sums else "-", rules_, lg["status"]])
    sheet("대사", ["모델", "원천 행", "규칙으로 제외", "모델 행", "합계 대사", "적용 규칙(행)", "상태"],
          data, [30, 10, 12, 10, 10, 40, 8])

    # ---- 원천정정요청
    data = []
    seen = set()
    for r_ in rules:
        rv, of = _parent_review(r_["id"], review)
        req = rv.get("request", "")
        if not req or (of, req) in seen:
            continue
        seen.add((of, req))
        n = sum(cnt.get(x["id"], 0) for x in rules if _parent_review(x["id"], review)[1] == of)
        data.append([of, r_.get("issue", ""), _source_of(cfg, r_["model"], manifest), req, n, "", "", ""])
    sheet("원천정정요청", ["규칙", "이슈", "원천 시스템·테이블", "요청 내용", "근거 행(처리행 시트)", "담당 부서", "요청일", "회신·조치"],
          data, [8, 8, 26, 50, 12, 14, 11, 30], inputs=(5, 6, 7))

    # ---- S2회귀
    sheet("S2회귀", ["ID", "대상", "원천 확인 결과", "선택지", "결정", "결정일"],
          [[x.get("id", ""), x.get("subject", ""), x.get("finding", ""), x.get("options", ""), x.get("decision", ""),
            x.get("decided", "")] for x in regress], [9, 18, 50, 36, 50, 11], inputs=(4, 5))

    # ---- 승인
    ws = sheet("승인", ["역할", "이름", "일자", "의견"],
               [[s.get("role", ""), s.get("name", ""), s.get("date", ""), s.get("note", "")] for s in appr.get("signers", [])],
               [30, 18, 12, 60], inputs=(1, 2, 3))
    ws.insert_rows(1, 4)
    ws["A1"] = "결정"
    ws["B1"] = APPROVAL_KO.get(appr.get("status", ""), appr.get("status", "대기")) if appr else "대기"
    ws["C1"] = appr.get("date", "")
    ws["D1"] = appr.get("decision", "")
    ws["A2"] = "다음 단계 과제"
    ws["D2"] = "\n".join(f"· {c}" for c in appr.get("carry_over", []))
    ws["D2"].alignment = wrap
    ws["A1"].font = ws["A2"].font = Font(bold=True)
    ws.freeze_panes = None

    for w in wb.worksheets:                                   # 인쇄: 가로, 폭 맞춤
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
        w.sheet_properties.pageSetUpPr.fitToPage = True

    path = out or os.path.join(cfg.mapping_dir, f"S3_정제규칙_검토서_{cfg.ent}.xlsx")
    wb.save(path)
    undecided = sum(1 for r_ in rules if not _parent_review(r_["id"], review)[0].get("decision"))
    print(f"ksxload review: 규칙 {len(rules)}개(미결정 {undecided}) · 처리 행 {len(q)} · 원천 정정 요청 "
          f"{wb['원천정정요청'].max_row - 1}건 → {os.path.relpath(path, cfg.repo)}")
    return path
