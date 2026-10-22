"""시나리오 정의서(docx)·업무 질의서(xlsx) 생성 — S1 절차 1~3.

입력: s1_config.yaml(시나리오·대상 명사·공백), scenarios.yaml(목적·흐름·질의서·사건), s1_extract_{ent}.json(원천 분석 결과)
"""
from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from datetime import date

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .output import _sheet, Workbook

FONT = "맑은 고딕"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def _font(run_or_style, size=None, bold=None, color=None):
    f = run_or_style.font
    f.name = FONT
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color
    el = run_or_style.element if hasattr(run_or_style, "element") else run_or_style._element
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        _font(r, 9, True)
        _shade(c, "DCE6F2")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = str(v if v is not None else "").split("\n")
            p = cells[i].paragraphs[0]
            for j, line in enumerate(lines):
                if j:
                    p = cells[i].add_paragraph()
                _font(p.add_run(line), 9)
    t.autofit = False
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 567)))          # cm → twip
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def _p(doc, text, size=10, bold=False, style=None):
    p = doc.add_paragraph(style=style)
    _font(p.add_run(text), size, bold)
    return p


def _h(doc, text, level):
    h = doc.add_heading(level=level)
    _font(h.add_run(text), {0: 20, 1: 15, 2: 12}[level], True, ACCENT)
    return h


def build(cfg_path, repo):
    with open(cfg_path, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    base = os.path.join(repo, cfg["output_dir"])
    with open(os.path.join(base, "scenarios.yaml"), encoding="utf-8") as f:
        sd = yaml.safe_load(f)
    ent = cfg["enterprise"]
    ext_path = os.path.join(base, f"s1_extract_{ent}.json")
    ext = json.load(open(ext_path, encoding="utf-8")) if os.path.exists(ext_path) else {"tables": [], "issues": []}
    scen = {s["id"]: s for s in cfg["scenarios"]}
    hints = cfg.get("noun_hints", [])
    rows_by_table = {f"{t['schema']}.{t['table']}": t["rows"] for t in ext["tables"]}
    issues_by_table = defaultdict(list)
    for i in ext.get("issues", []):
        issues_by_table[i["table"]].append(i)
    qs = sd["questions"]

    # ---------------- 업무 질의서 xlsx ----------------
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "업무질의서", ["ID", "시나리오", "유형", "질문", "판단 기준", "필요 명사", "기대 답(실습 데이터)", "S7 규칙 후보", "S9 평가셋"],
           [(q["id"], f"{q['sc']} {scen[q['sc']]['name']}", q["type"], q["q"], q["criteria"], ", ".join(q["nouns"]),
             q.get("answer", ""), "Y" if q["type"] in ("판단", "진단", "추적") else "", "Y")
            for q in qs], [11, 18, 6, 44, 34, 30, 44, 8, 7], wrap_cols=(4, 5, 6, 7))
    cnt = Counter((q["sc"], q["type"]) for q in qs)
    types = sd.get("types", sorted({q["type"] for q in qs}))
    _sheet(wb, "시나리오별 집계", ["시나리오"] + types + ["합계", "기대 답 보유"],
           [[f"{sid} {scen[sid]['name']}"] + [cnt[(sid, t)] for t in types] +
            [sum(cnt[(sid, t)] for t in types), sum(1 for q in qs if q["sc"] == sid and q.get("answer"))]
            for sid in scen], [24] + [7] * len(types) + [7, 10])
    _sheet(wb, "명사별 사용", ["명사", "질의 수", "시나리오"],
           sorted(((n, sum(1 for q in qs if n in q["nouns"]),
                    ", ".join(sorted({q["sc"] for q in qs if n in q["nouns"]}))) for n in {n for q in qs for n in q["nouns"]}),
                  key=lambda r: -r[1]), [30, 8, 30])
    _sheet(wb, "사건(정답지)", ["ID", "사건", "기간", "내용"],
           [(e["id"], e["name"], e["when"], e["detail"]) for e in sd["events"]], [6, 26, 26, 70], wrap_cols=(4,))
    xlsx = os.path.join(base, f"S1_업무질의서_{ent}.xlsx")
    wb.save(xlsx)

    # ---------------- 시나리오 정의서 docx ----------------
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)
    _font(doc.styles["Normal"], 10)
    c = sd["company"]
    t = doc.add_paragraph()
    _font(t.add_run(f"{ent} 1차 업무 시나리오 정의서"), 22, True, ACCENT)
    _p(doc, "S1 원천 분석 산출물 · KS X 9101-1:2025 개정(안) v1.0 기준", 11)
    _table(doc, ["항목", "내용"], [
        ("대상 기업", c["name"]), ("작성일", date.today().isoformat()), ("상태", "초안 — 현업 승인 전"),
        ("관련 산출물", f"S1_원천인벤토리_{ent}.xlsx, S1_업무질의서_{ent}.xlsx, scenarios.yaml"),
        ("완료 기준", c["approval"])], [4, 13])

    _h(doc, "1. 기업 개요", 1)
    _table(doc, ["항목", "내용"], [
        ("제품", c["product"]), ("사업장", c["site"]), ("공장", "\n".join(c["plants"])),
        ("공정 흐름", c["flow"]), ("원천 시스템", "\n".join(c["systems"])), ("데이터 기간", c["period"])], [3, 14])

    _h(doc, "2. 시나리오 요약", 1)
    _p(doc, f"1차 대상은 구현안내서 4.1 권장 5종이다. 업무 질의는 시나리오별 20건, 총 {len(qs)}건이다.")
    summ = []
    for sid, s in scen.items():
        tabs = [h["table"] for h in hints if sid in h.get("scenarios", [])]
        summ.append((sid, s["name"], sd["scenarios"][sid]["purpose"], len(s["nouns"]), len(tabs),
                     sum(1 for q in qs if q["sc"] == sid)))
    _table(doc, ["ID", "시나리오", "목적", "대상 명사", "원천 테이블", "질의"], summ, [1.3, 2.8, 8.5, 1.5, 1.6, 1.3])

    for n, (sid, s) in enumerate(scen.items(), 3):
        d = sd["scenarios"][sid]
        _h(doc, f"{n}. {sid} {s['name']}", 1)
        _p(doc, d["purpose"])
        _h(doc, "업무 흐름", 2)
        for i, step in enumerate(d["flow"], 1):
            _p(doc, f"{i}) {step}")
        _h(doc, "범위와 KPI", 2)
        _table(doc, ["항목", "내용"], [("범위", d["scope"]), ("KPI", ", ".join(d["kpi"]))], [3, 14])
        _h(doc, "대상 명사와 원천 대응(1차)", 2)
        rows = []
        for noun in s["nouns"] + s.get("ref_nouns", []):
            src = [f"{h['table']} ({rows_by_table.get(h['table'], 0):,}행)" for h in hints
                   if noun in h.get("nouns", []) and h.get("scope", "대상") != "제외(ERP 원장 사용)"]
            rows.append((noun, "대상" if noun in s["nouns"] else "참조", "\n".join(src) or "원천 없음"))
        _table(doc, ["표준 명사", "구분", "원천 테이블(행 수)"], rows, [5, 1.8, 10.2])
        _h(doc, "원천 공백·Extension 후보", 2)
        for g in s.get("gap", "").split(". "):
            if g.strip():
                _p(doc, "• " + g.strip().rstrip("."))
        tabs = [h["table"] for h in hints if sid in h.get("scenarios", [])]
        iss = sorted((i for tb in tabs for i in issues_by_table.get(tb, []) if i["severity"] in ("상", "중")),
                     key=lambda i: i["id"])
        if iss:
            _h(doc, "관련 품질 이슈(상·중)", 2)
            _table(doc, ["ID", "심각도", "위치", "이슈", "건수"],
                   [(i["id"], i["severity"], f"{i['table']}.{i['column']}".rstrip("."), i["issue"], i["count"])
                    for i in iss], [1.4, 1.3, 5.3, 7.6, 1.4])
        _h(doc, "업무 질의(20건)", 2)
        _table(doc, ["ID", "유형", "질문", "판단 기준"],
               [(q["id"].replace("Q-", ""), q["type"], q["q"], q["criteria"]) for q in qs if q["sc"] == sid],
               [2.0, 1.2, 7.8, 6.0])

    _h(doc, f"{len(scen) + 3}. 승인", 1)
    _p(doc, "S1 완료 기준: 대상 테이블 100% 프로파일, 시나리오별 대상 명사 목록 승인(구현안내서 표 7).")
    _table(doc, ["역할", "이름", "승인일", "의견"],
           [("도메인(PI) 컨설턴트", "", "", ""), ("고객사 현업 검수자", "", "", ""), ("PM", "", "", "")], [4.5, 3.5, 3, 6])

    _h(doc, "부록 A. 실습 데이터 사건(정답지)", 1)
    _p(doc, "가상 원천 데이터에 의도적으로 심은 사건이다. S7 추론 규칙과 S9 AI 평가의 정답으로 쓴다.")
    _table(doc, ["ID", "사건", "기간", "내용"], [(e["id"], e["name"], e["when"], e["detail"]) for e in sd["events"]],
           [1.2, 3.8, 3.8, 8.2])
    docx = os.path.join(base, f"S1_시나리오정의서_{ent}.docx")
    doc.save(docx)
    return xlsx, docx
