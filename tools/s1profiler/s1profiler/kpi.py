"""KPI 정의서 생성 + 원천 산출 검증 (S1 확장).

kpi_catalog.yaml 의 datasets(원천별 추출 SQL) → 메모리 sqlite 적재 → views → KPI별 result SQL 순으로 계산하고
S1_KPI정의서_{ent}.xlsx 와 s1_kpi_{ent}.json 을 만든다.

  python -m s1profiler kpi --config mapping/impix/s1/s1_config.yaml [--catalog mapping/impix/s1/kpi_catalog.yaml]
"""
from __future__ import annotations

import json
import math
import os
import sqlite3
from collections import OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import core
from .output import BORDER, FONT, HEAD_FILL, _sheet

AREA_FILL = {"생산관리": "E8F1FB", "품질관리": "FDF0E6", "SCM": "EAF6EC", "고객서비스": "F3EDFA"}


def _d8(v):
    if v is None:
        return None
    s = str(v)
    return f"{s[:4]}-{s[4:6]}-{s[6:8]}" if len(s) >= 8 and s[:8].isdigit() else s


class _Prod:
    def __init__(self):
        self.v, self.n = 1.0, 0

    def step(self, x):
        if x is not None:
            self.v *= float(x)
            self.n += 1

    def finalize(self):
        return self.v if self.n else None


def _norm(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, int):          # psql 대체 드라이버는 정수값 float8을 int로 돌려준다 → sqlite 정수 나눗셈 방지
        return float(v)
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    return v


def _subst(sql, params):
    for k in sorted(params, key=len, reverse=True):
        sql = sql.replace(":" + k, "'" + str(params[k]).replace("'", "''") + "'")
    return sql


def compute(cfg, cat, conns=None, log=print):
    """KPI 계산. 반환: (datasets 행 수, KPI 결과 dict, 일별 dict)."""
    as_of = cat["as_of"]
    params = {"as_of": as_of, "as_of_date": as_of[:10], "as_of_ym": as_of[:7]}
    own = conns is None
    if own:
        conns = {s["system"]: core.connect(s["dsn"]) for s in cfg["sources"]}
    db = sqlite3.connect(":memory:")
    db.create_function("d8", 1, _d8)
    db.create_aggregate("prod", 1, _Prod)
    # 달력
    c = cat["calendar"]
    d0, d1 = date.fromisoformat(c["from"]), date.fromisoformat(c["to"])
    off = set(c.get("off_weekdays", [7]))
    db.execute("CREATE TABLE cal (d TEXT PRIMARY KEY, ym TEXT, is_work INTEGER)")
    db.executemany("INSERT INTO cal VALUES (?,?,?)",
                   [((d0 + timedelta(i)).isoformat(), (d0 + timedelta(i)).isoformat()[:7],
                     0 if (d0 + timedelta(i)).isoweekday() in off else 1) for i in range((d1 - d0).days + 1)])
    db.execute("CREATE TABLE months (ym TEXT)")
    db.executemany("INSERT INTO months VALUES (?)", [(m,) for m in cat["months"]])
    counts = OrderedDict()
    for name, ds in cat["datasets"].items():
        cur = conns[ds["system"]].execute(ds["sql"].replace("{as_of}", as_of))
        rows = [tuple(_norm(v) for v in r) for r in cur.fetchall()]
        cols = [getattr(d, "name", None) or d[0] for d in (cur.description or [])]
        if not cols:
            raise RuntimeError(f"데이터셋 {name}: 결과 열을 알 수 없음(행 0건)")
        db.execute(f"CREATE TABLE {name} ({', '.join(cols)})")
        db.executemany(f"INSERT INTO {name} VALUES ({','.join('?' * len(cols))})", rows)
        counts[name] = (ds["system"], len(rows))
    log(f"[KPI] 원천 추출 {len(counts)}개 데이터셋, {sum(n for _, n in counts.values()):,}행")
    for v in cat.get("views", []):
        db.execute(_subst(v, params))
    results, daily = OrderedDict(), OrderedDict()
    for k in cat["kpis"]:
        cur = db.execute(_subst(k["result"], params))
        names = [d[0] for d in cur.description]
        rows = []
        for r in cur.fetchall():
            m = dict(zip(names, r))
            rows.append({"period": m.get("period"), "dim": m.get("dim") or "", "value": m.get("value"),
                         "num": m.get("num"), "den": m.get("den"), "note": m.get("note") or ""})
        db.execute(f"CREATE TABLE k_{k['id']} (period, dim, value, num, den, note)")
        db.executemany(f"INSERT INTO k_{k['id']} VALUES (?,?,?,?,?,?)",
                       [(r["period"], r["dim"], r["value"], r["num"], r["den"], r["note"]) for r in rows])
        results[k["id"]] = rows
        if k.get("daily"):
            daily[k["id"]] = [dict(zip(("d", "dim", "value"), r)) for r in db.execute(_subst(k["daily"], params))]
    log(f"[KPI] {len(results)}개 KPI 계산, 결과 {sum(len(v) for v in results.values())}행")
    if own:
        for cn in conns.values():
            cn.close()
    return counts, results, daily


def _fmt(v, unit):
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if unit in ("원", "EA", "건", "PPM"):
        return round(v)
    return round(v, 2)


def _headline(rows, period, key=None):
    """KPI 대표값: catalog headline 차원 → '전체'·'합계' 행 → dim 1개일 때 그 값. 시점 KPI는 같은 월의 시점값."""
    rs = [r for r in rows if r["period"] == period] or [r for r in rows if str(r["period"]).startswith(period)]
    for key in ([key] if key else []) + ["전체", "합계", "5 합계", "실패비용 합계", "FPY(직행률)", "1 180일 초과(장기 체화)", "품목 일치율"]:
        for r in rs:
            if r["dim"] == key:
                return r["value"]
    return rs[0]["value"] if len(rs) == 1 else None


def write(out_dir, cfg, cat, counts, results, daily, questions=None):
    ent = cfg["enterprise"]
    os.makedirs(out_dir, exist_ok=True)
    kpis = cat["kpis"]
    qmap = defaultdict(list)
    for q in questions or []:
        for kid in q.get("kpi", []):
            qmap[kid].append(q["id"])
    months = cat["months"]
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    wb = Workbook()
    wb.remove(wb.active)

    # 1. 개요
    by_area = OrderedDict()
    for k in kpis:
        by_area.setdefault(f"{k['area']} — {k['group']}", []).append(k["id"])
    ov = [("대상 기업", ent), ("생성 일시", stamp), ("기준 일시(as_of)", cat["as_of"]), ("평가 월", ", ".join(months)),
          ("데이터 기간", cat.get("period_note", "")), ("KPI 수", len(kpis)),
          ("원천 데이터셋", ", ".join(f"{n}({s} {c:,}행)" for n, (s, c) in counts.items())),
          ("산출 방식", "원천 SQL(PostgreSQL) → 메모리 sqlite → KPI SQL. 시트 '산출SQL' 참조"),
          ("정의 원본", "mapping/impix/s1/kpi_catalog.yaml (수정 후 s1profiler kpi 재실행)")]
    ov += [(a, ", ".join(ids)) for a, ids in by_area.items()]
    _sheet(wb, "개요", ["항목", "내용"], ov, [26, 110], wrap_cols=(2,))

    # 2. KPI 정의서
    hdr = ["ID", "영역", "그룹", "KPI명", "정의", "산식", "단위", "주기", "방향", "목표(예시)", "분석 차원",
           "표준 명사(KS X 9101)", "표준 속성", "원천(시스템·테이블)", "원천 공백·Extension", "시나리오", "관련 질의"]
    hdr += ["대표 차원"] + [f"{m} 대표값" for m in months]
    rows = []
    for k in kpis:
        r = [k["id"], k["area"], k["group"], k["name"], k["definition"], k["formula"], k["unit"], k["period"],
             k["direction"], k.get("target", ""), ", ".join(k.get("dims", [])), ", ".join(k["nouns"]),
             "\n".join(k.get("attrs", [])), k["source"], k.get("gap", ""), ", ".join(k.get("scenarios", [])),
             ", ".join(qmap.get(k["id"], [])), k.get("headline", "전체")]
        r += [_fmt(_headline(results[k["id"]], m, k.get("headline")), k["unit"]) for m in months]
        rows.append(r)
    ws = _sheet(wb, "KPI정의서", hdr, rows,
                [6, 9, 12, 20, 40, 36, 6, 8, 10, 14, 16, 26, 34, 26, 36, 9, 18, 14] + [11] * len(months),
                wrap_cols=(4, 5, 6, 10, 11, 12, 13, 14, 15, 17))
    for row in ws.iter_rows(min_row=2):
        f = AREA_FILL.get(row[1].value)
        if f:
            for cell in row[:4]:
                cell.fill = PatternFill("solid", fgColor=f)

    # 3. KPI 산출값
    val_rows = []
    unit = {k["id"]: k["unit"] for k in kpis}
    for k in kpis:
        for r in results[k["id"]]:
            val_rows.append((k["id"], k["name"], r["period"], r["dim"], _fmt(r["value"], k["unit"]), k["unit"],
                             k.get("num_label", "분자") if r["num"] is not None else "",
                             round(r["num"], 2) if isinstance(r["num"], float) else r["num"],
                             k.get("den_label", "분모") if r["den"] is not None else "",
                             round(r["den"], 2) if isinstance(r["den"], float) else r["den"], r["note"]))
    _sheet(wb, "KPI산출값", ["ID", "KPI명", "기간", "차원", "값", "단위", "분자 항목", "분자", "분모 항목", "분모", "비고(근거)"],
           val_rows, [6, 22, 11, 26, 12, 6, 16, 14, 14, 14, 60], wrap_cols=(11,))

    # 4. 월별 비교(피벗)
    piv = []
    for k in kpis:
        dims = OrderedDict()
        for r in results[k["id"]]:
            dims.setdefault(r["dim"], {})[r["period"]] = r["value"]
        for dname, per in dims.items():
            line = [k["id"], k["name"], dname, k["unit"]]
            if all(p not in months for p in per):          # 시점 KPI
                line += [None] * (len(months) - 1) + [_fmt(next(iter(per.values())), k["unit"])]
            else:
                line += [_fmt(per.get(m), k["unit"]) for m in months]
            piv.append(line)
    _sheet(wb, "월별비교", ["ID", "KPI명", "차원", "단위"] + months, piv, [6, 24, 28, 7] + [13] * len(months))

    # 5. 일별 추이
    drows = [(kid, next(k["name"] for k in kpis if k["id"] == kid), r["d"], r["dim"], _fmt(r["value"], unit[kid]))
             for kid, rs in daily.items() for r in rs]
    _sheet(wb, "일별추이", ["ID", "KPI명", "일자", "차원", "값"], drows, [6, 24, 11, 22, 12])

    # 6. 명사-KPI 매트릭스
    nouns = sorted({n for k in kpis for n in k["nouns"]})
    mrows = [[n] + ["●" if n in k["nouns"] else "" for k in kpis] + [sum(1 for k in kpis if n in k["nouns"])] for n in nouns]
    ws = _sheet(wb, "명사-KPI", ["표준 명사"] + [k["id"] for k in kpis] + ["KPI 수"], mrows, [30] + [5] * len(kpis) + [7])
    for row in ws.iter_rows(min_row=2):
        for cell in row[1:-1]:
            cell.alignment = Alignment(horizontal="center")

    # 7. 공백·Extension
    grows = [(k["id"], k["name"], k["gap"], "Extension" if "Extension" in k["gap"] else ("Core" if k["gap"].startswith("없음") else "검토"))
             for k in kpis]
    _sheet(wb, "공백·Extension", ["ID", "KPI명", "원천 공백·Extension 제안", "구분"], grows, [6, 26, 90, 10], wrap_cols=(3,))

    # 8. 산출 SQL
    srows = [("[데이터셋] " + n, cat["datasets"][n]["system"], cat["datasets"][n].get("desc", ""), cat["datasets"][n]["sql"].strip())
             for n in cat["datasets"]]
    srows += [(f"[뷰 {i}]", "sqlite", "", v.strip()) for i, v in enumerate(cat.get("views", []), 1)]
    for k in kpis:
        srows.append((k["id"], "sqlite", k["name"], k["result"].strip()))
        if k.get("daily"):
            srows.append((k["id"] + " 일별", "sqlite", k["name"], k["daily"].strip()))
    ws = _sheet(wb, "산출SQL", ["대상", "엔진", "설명", "SQL"], srows, [16, 8, 30, 120], wrap_cols=(4,))
    for row in ws.iter_rows(min_row=2):
        row[3].font = Font(name="Consolas", size=8)

    xlsx = os.path.join(out_dir, f"S1_KPI정의서_{ent}.xlsx")
    wb.save(xlsx)
    js = {"enterprise": ent, "as_of": cat["as_of"], "generated": stamp,
          "kpis": [{**{f: k.get(f) for f in ("id", "area", "group", "name", "unit", "nouns", "scenarios")},
                    "questions": qmap.get(k["id"], []), "values": results[k["id"]]} for k in kpis]}
    jpath = os.path.join(out_dir, f"s1_kpi_{ent}.json")
    with open(jpath, "w", encoding="utf-8") as f:
        json.dump(js, f, ensure_ascii=False, indent=1, default=str)
    return xlsx, jpath


def run(cfg, cat_path, repo, log=print):
    with open(cat_path, encoding="utf-8") as f:
        cat = yaml.safe_load(f)
    counts, results, daily = compute(cfg, cat, log=log)
    out_dir = os.path.join(repo, cfg["output_dir"])
    qs = []
    sp = os.path.join(out_dir, "scenarios.yaml")
    if os.path.exists(sp):
        with open(sp, encoding="utf-8") as f:
            qs = yaml.safe_load(f).get("questions", [])
    return write(out_dir, cfg, cat, counts, results, daily, qs)
