"""S1 산출물 기록: 인벤토리 xlsx(시트 9종), 요약 md, S2 입력용 json."""
from __future__ import annotations

import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="DCE6F2")
SEV_FILL = {"상": PatternFill("solid", fgColor="F8D7DA"), "중": PatternFill("solid", fgColor="FFF3CD"),
            "하": PatternFill("solid", fgColor="EEF2F7")}
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "맑은 고딕"


def _sheet(wb, title, header, rows, widths=None, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append([_cell(v) for v in r])
    for c in range(1, len(header) + 1):
        h = ws.cell(row=1, column=c)
        h.font = Font(name=FONT, bold=True, size=10)
        h.fill = HEAD_FILL
        h.alignment = Alignment(vertical="center", wrap_text=True)
        h.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=9)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_cols)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _cell(v):
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, float):
        return round(v, 4)
    return v


def write_all(out_dir, cfg, tables, inventory, masters, codes, relations, issues):
    os.makedirs(out_dir, exist_ok=True)
    ent = cfg["enterprise"]
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    wb = Workbook()
    wb.remove(wb.active)

    # 1. 요약
    n_cols = sum(len(t.columns) for t in tables)
    n_cmt = sum(1 for t in tables for c in t.columns if c.comment)
    sev = {s: sum(1 for i in issues if i["severity"] == s) for s in ("상", "중", "하")}
    summary = [
        ("대상 기업", ent), ("생성 일시", stamp), ("원천 시스템", ", ".join(s["system"] for s in cfg["sources"])),
        ("테이블 수", len(tables)), ("1차 범위 테이블", sum(1 for i in inventory if i["scope"] == "대상")),
        ("컬럼 수", n_cols), ("컬럼 주석 비율(%)", round(100.0 * n_cmt / n_cols, 1) if n_cols else 0),
        ("총 행 수", sum(t.rows for t in tables)),
        ("선언 FK", sum(1 for r in relations if r["declared"])),
        ("암묵 관계(FK 미선언)", sum(1 for r in relations if not r["declared"])),
        ("시스템 간 관계", sum(1 for r in relations if r.get("cross_system"))),
        ("코드 마스터 코드 수", len(masters)), ("코드 후보 컬럼", len(codes)),
        ("품질 이슈(상/중/하)", f"{sev['상']} / {sev['중']} / {sev['하']}"),
        ("프로파일 완료율", "100% (전 테이블 전수 집계)"),
    ]
    _sheet(wb, "요약", ["항목", "값"], summary, [28, 60])

    # 2. 원천 인벤토리(구현안내서 표 8 + 보조 열)
    _sheet(wb, "원천인벤토리",
           ["시스템", "스키마", "테이블", "업무 의미", "행 수", "일 증가량(활동일 평균)", "증가량 기준", "키 후보",
            "대응 예상 명사", "시나리오", "1차 범위", "담당자", "컬럼 수", "주석률(%)", "선언 FK", "암묵 관계", "품질 이슈", "비고"],
           [(i["system"], i["schema"], i["table"], i["meaning"], i["rows"], i["per_day"], i["daily_basis"], i["keys"],
             i["nouns"], i["scenarios"], i["scope"], i["owner"], i["columns"], i["comment_cov"], i["fks"],
             i["implicit"], i["issues"], i["note"]) for i in inventory],
           [8, 7, 16, 30, 9, 11, 26, 28, 30, 14, 14, 18, 7, 8, 7, 8, 8, 30], wrap_cols=(4, 8, 9, 18))

    # 3. 컬럼 목록 + 4. 프로파일
    col_rows, prof_rows = [], []
    for t in tables:
        for c in t.columns:
            p = t.profile.get(c.name, {})
            col_rows.append((t.system, t.schema, t.name, c.ordinal, c.name, c.type, "N" if c.not_null else "Y",
                             "Y" if c.pk else "", c.fk or "", c.comment or "", c.default or ""))
            dl = p.get("date_like")
            prof_rows.append((t.system, t.name, c.name, c.type, t.rows, p.get("null", 0), p.get("null_pct", 0),
                              p.get("distinct", 0), p.get("min"), p.get("max"), p.get("avg_len"),
                              round(p["mean"], 3) if p.get("mean") is not None else None,
                              round(p["stddev"], 3) if p.get("stddev") is not None else None,
                              "; ".join(f"{v}({n})" for v, n in (p.get("top") or [])),
                              "; ".join(f"{v}({n})" for v, n in (p.get("patterns") or [])),
                              "/".join(dl["formats"]) if dl else "", "Y" if p.get("key_candidate") else "",
                              p.get("whitespace") or ""))
    _sheet(wb, "컬럼목록", ["시스템", "스키마", "테이블", "순번", "컬럼", "타입", "NULL 허용", "PK", "FK(선언)", "주석", "기본값"],
           col_rows, [8, 7, 16, 6, 18, 22, 8, 5, 24, 28, 20])
    _sheet(wb, "컬럼프로파일",
           ["시스템", "테이블", "컬럼", "타입", "행 수", "NULL 수", "NULL(%)", "고유값 수", "최소", "최대", "평균 길이",
            "평균", "표준편차", "상위 값(건수)", "값 패턴(건수)", "문자열 일시 형식", "키 후보", "앞뒤 공백 수"],
           prof_rows, [8, 16, 16, 18, 8, 8, 8, 9, 16, 16, 8, 9, 9, 40, 30, 16, 7, 8], wrap_cols=(14, 15))

    # 5. 코드 마스터, 6. 코드 후보
    _sheet(wb, "코드마스터", ["시스템", "테이블", "코드그룹", "코드", "코드명"],
           [(m["system"], m["table"], m["group"], m["code"], m["name"]) for m in masters], [8, 16, 16, 12, 24])
    _sheet(wb, "코드후보컬럼",
           ["시스템", "테이블", "컬럼", "주석", "고유값 수", "값(건수)", "판정", "연결 코드그룹", "코드마스터에 없는 값", "주석 정의 코드"],
           [(c["system"], c["table"], c["column"], c["comment"] or "", c["distinct"],
             "; ".join(f"{k}({v})" for k, v in c["values"].items()), c["kind"], c["link"],
             ", ".join(c["unmatched"]), "; ".join(f"{k}={v}" for k, v in c["comment_codes"].items()))
            for c in codes], [8, 18, 14, 26, 8, 44, 22, 16, 20, 30], wrap_cols=(6, 10))

    # 7. 관계
    _sheet(wb, "관계(FK·암묵)",
           ["자식 테이블", "자식 컬럼", "부모 테이블", "부모 컬럼", "부모 조건", "발견 경로", "근거 SQL", "선언 FK",
            "시스템 간", "자식 고유값", "일치", "포함률", "고아 수", "고아 예", "비고"],
           [(r["child"], r["child_col"], r["parent"], r["parent_col"],
             ", ".join(f"{k}='{v}'" for k, v in r["parent_filter"].items()), ", ".join(r["found_by"]),
             ", ".join(r["queries"]), "Y" if r["declared"] else "", "Y" if r.get("cross_system") else "",
             r["child_distinct"], r["matched"], r["inclusion"], r["orphans"], ", ".join(r["orphan_samples"]),
             r.get("note") or "") for r in relations],
           [18, 14, 18, 14, 18, 18, 18, 7, 7, 8, 7, 7, 7, 24, 34], wrap_cols=(15,))

    # 8. 품질 이슈
    ws = _sheet(wb, "품질이슈", ["ID", "심각도", "시스템", "테이블", "컬럼", "이슈", "건수", "예시", "조치 제안"],
                [(i["id"], i["severity"], i["system"], i["table"], i["column"], i["issue"], i["count"], i["sample"],
                  i["action"]) for i in issues], [7, 6, 7, 18, 14, 30, 7, 40, 36], wrap_cols=(6, 8, 9))
    for row in ws.iter_rows(min_row=2):
        f = SEV_FILL.get(row[1].value)
        if f:
            row[1].fill = f

    # 9. 시나리오별 대상 명사
    sc_rows = []
    for s in cfg.get("scenarios", []):
        tabs = [i["fqn"] for i in inventory if s["id"] in i["scenarios"].split(", ")]
        sc_rows.append((s["id"], s["name"], ", ".join(s["nouns"]), ", ".join(s.get("ref_nouns", [])), "\n".join(tabs),
                        s.get("gap", "")))
    _sheet(wb, "시나리오-명사", ["ID", "시나리오", "대상 명사", "참조 명사(포함)", "원천 테이블", "원천 공백·확장 후보"],
           sc_rows, [7, 20, 44, 40, 30, 40], wrap_cols=(3, 4, 5, 6))

    xlsx = os.path.join(out_dir, f"S1_원천인벤토리_{ent}.xlsx")
    wb.save(xlsx)

    # JSON(S2 automap 입력)
    js = {"enterprise": ent, "generated": stamp,
          "tables": [{"system": t.system, "db": t.db, "schema": t.schema, "table": t.name, "comment": t.comment,
                      "rows": t.rows, "pk": t.pk, "fks": t.fks, "daily": t.daily,
                      "columns": [{"name": c.name, "type": c.type, "not_null": c.not_null, "comment": c.comment,
                                   "pk": c.pk, "fk": c.fk, "profile": t.profile.get(c.name)} for c in t.columns]}
                     for t in tables],
          "codes": {"masters": masters, "candidates": codes}, "relations": relations, "issues": issues}
    jpath = os.path.join(out_dir, f"s1_extract_{ent}.json")
    with open(jpath, "w", encoding="utf-8") as f:
        json.dump(js, f, ensure_ascii=False, indent=1, default=str)

    # 요약 MD
    md = [f"# S1 원천 분석 요약 — {ent}", "", f"생성: {stamp} · 도구: tools/s1profiler", "",
          "| 항목 | 값 |", "| --- | --- |"] + [f"| {k} | {v} |" for k, v in summary] + [""]
    md += ["## 품질 이슈 (상)", "", "| ID | 테이블.컬럼 | 이슈 | 건수 | 예시 |", "| --- | --- | --- | --- | --- |"]
    md += [f"| {i['id']} | {i['table']}.{i['column']} | {i['issue']} | {i['count']} | {i['sample'][:60]} |"
           for i in issues if i["severity"] == "상"]
    md += ["", "## 암묵 관계 (FK 미선언)", "", "| 자식 | 부모 | 발견 경로 | 포함률 | 고아 |", "| --- | --- | --- | --- | --- |"]
    md += [f"| {r['child']}.{r['child_col']} | {r['parent']}.{r['parent_col']} | {', '.join(r['found_by'])} | "
           f"{r['inclusion']} | {r['orphans']} |" for r in relations if not r["declared"]]
    mpath = os.path.join(out_dir, f"S1_원천분석_요약_{ent}.md")
    with open(mpath, "w", encoding="utf-8") as f:
        f.write("\n".join(md) + "\n")
    return xlsx, jpath, mpath
